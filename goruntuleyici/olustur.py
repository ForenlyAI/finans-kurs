"""Yerel görüntüleyici: GERÇEK .xlsx dosyasını (hücreler, biçimler, formüller) + hesaplanmış değerleri HTML'e çevirir.

Kullanım (lab klasöründen):
    uv run python goruntuleyici/olustur.py        → goruntuleyici/site/index.html + <kitap>.html

Sayfa adresi seçenekleri (tarayıcıda ya da ekran görüntüsünde):
    <kitap>.html?sayfa=2              2. sayfa sekmesi (1'den başlar)
    <kitap>.html?sayfa=2&hucre=C8     C8 seçili, formül çubuğunda formülü
    <kitap>.html?sayfa=2&formul=1     hücrelerde değer yerine formül metni (Excel/Sheets "formülleri göster")
    <kitap>.html?sayfa=2&satir=1-14,30-37   yalnız bu satırlar (ekrana sığdırmak için; satır numaraları görünür kalır)
    <kitap>.html?sayfa=2&olcek=0.9    tabloyu yakınlaştır/uzaklaştır (geniş 12 aylık tablolar için)

Bu bir Excel ya da Google E-Tablolar ekranı DEĞİLDİR ve öyle görünmeye çalışmaz: kursun kendi
görüntüleyicisidir. Değerler kitaplar/hesapla.py'nin formül motoru çıktısıdır (hesap/<kitap>.json).
"""

from __future__ import annotations

import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

LAB = Path(__file__).resolve().parent.parent
SITE = Path(__file__).resolve().parent / "site"

# Türkçe Excel işlev adları — Microsoft Destek tr-tr sayfalarından (KAYNAKLAR.yml SRC-13), 2026-09-25
TR_ISLEV = {"SUMIF": "ETOPLA", "COUNTIF": "EĞERSAY", "SUM": "TOPLA", "IFERROR": "EĞERHATA", "IF": "EĞER",
            "ROUNDUP": "YUKARIYUVARLA", "ROUND": "YUVARLA", "MIN": "MİN", "MAX": "MAK", "AVERAGE": "ORTALAMA",
            "STDEV.P": "STDSAPMA.P", "PMT": "DEVRESEL_ÖDEME", "IPMT": "FAİZTUTARI", "FV": "GD", "AND": "VE"}


def tr_formul(f: str) -> str:
    def degis(m):
        return TR_ISLEV.get(m.group(1).upper(), m.group(1)) + "("
    # tırnak içindeki metne dokunma
    parcalar = re.split(r'("[^"]*")', f)
    out = []
    for p in parcalar:
        if p.startswith('"'):
            out.append(p)
        else:
            p = re.sub(r"([A-Z][A-Z0-9.]*)\(", degis, p)
            out.append(p.replace(",", ";"))
    return "".join(out)


def sayi_tr(x: float, ondalik: int, binlik: bool = True) -> str:
    s = f"{abs(x):,.{ondalik}f}"
    s = s.replace(",", "X").replace(".", ",").replace("X", "." if binlik else "")
    return ("−" if x < 0 and float(f"{abs(x):.{ondalik}f}") != 0 else "") + s


def bicimle(v, bicim: str | None) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "DOĞRU" if v else "YANLIŞ"
    if isinstance(v, str):
        return v
    b = bicim or "General"
    if "%" in b:
        ond = len(b.split(".")[1].rstrip("%")) if "." in b else 0
        s = sayi_tr(v * 100, ond)
        return ("−%" + s[1:]) if s.startswith("−") else "%" + s
    ond = len(re.search(r"\.(0+)", b).group(1)) if re.search(r"\.(0+)", b) else 0
    if '"TL"' in b:
        return sayi_tr(v, ond) + " TL"
    if b in ("General",):
        if float(v).is_integer():
            return sayi_tr(v, 0, binlik=abs(v) >= 10000)
        return sayi_tr(v, 4).rstrip("0").rstrip(",")
    return sayi_tr(v, ond, binlik="," in b)


def renk(c) -> str | None:
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and c.rgb not in ("00000000",):
            return "#" + c.rgb[-6:]
    except Exception:
        pass
    return None


def sayfa_html(ws, degerler: dict, idx: int) -> str:
    max_c = ws.max_column
    max_r = ws.max_row
    sutunlar = [get_column_letter(i) for i in range(1, max_c + 1)]
    colgroup = ['<col style="width:40px">']
    toplam = 40
    for s in sutunlar:
        w = ws.column_dimensions[s].width if s in ws.column_dimensions and ws.column_dimensions[s].width else 8.43
        px = round(w * 8.4 + 6)
        toplam += px
        colgroup.append(f'<col style="width:{px}px">')
    satirlar = ['<tr class="bas"><th></th>' + "".join(f"<th>{s}</th>" for s in sutunlar) + "</tr>"]
    for r in range(1, max_r + 1):
        hucreler = [f'<th class="no">{r}</th>']
        for s in sutunlar:
            c = ws[f"{s}{r}"]
            d = degerler.get(c.coordinate)
            f = d["f"] if d else None
            v = d["v"] if d else None
            gor = bicimle(v, c.number_format)
            stil = []
            dolgu = renk(c.fill.fgColor) if c.fill and c.fill.fill_type == "solid" else None
            if dolgu:
                stil.append(f"background:{dolgu}")
            yr = renk(c.font.color) if c.font and c.font.color else None
            if yr:
                stil.append(f"color:{yr}")
            if c.font and c.font.b:
                stil.append("font-weight:700")
            if c.font and c.font.i:
                stil.append("font-style:italic")
            if c.font and c.font.sz and c.font.sz != 11:
                stil.append(f"font-size:{c.font.sz / 11:.2f}em")
            sayisal = isinstance(v, (int, float)) and not isinstance(v, bool)
            sinif = ["h"]
            if sayisal:
                sinif.append("sag")
            if c.alignment and c.alignment.horizontal == "center":
                sinif.append("orta")
            if c.alignment and c.alignment.wrap_text:
                sinif.append("sar")
            if f:
                sinif.append("fx")
            if sayisal and v < 0:
                sinif.append("eksi")
            # metin, sağdaki boş hücrelere taşabilsin (Excel gibi) — yalnız uzun etiketlerde
            ozn = f' data-a="{c.coordinate}"'
            if f:
                ozn += f' data-f="{html.escape(f)}" data-tr="{html.escape(tr_formul(f))}"'
            ozn += f' data-v="{html.escape(gor)}"'
            hucreler.append(f'<td class="{" ".join(sinif)}"{ozn} style="{";".join(stil)}"><span>{html.escape(gor)}</span></td>')
        yuk = ws.row_dimensions[r].height
        satirlar.append(f'<tr data-r="{r}"' + (f' style="height:{round(yuk * 1.33)}px"' if yuk else "") + ">" + "".join(hucreler) + "</tr>")
    return (f'<section class="sayfa" data-i="{idx}" data-ad="{html.escape(ws.title)}">'
            f'<table style="width:{toplam}px"><colgroup>{"".join(colgroup)}</colgroup>{"".join(satirlar)}</table></section>')


SABLON = """<!DOCTYPE html>
<html lang="tr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{kitap} · Finans 101 lab görüntüleyici</title>
<style>
:root {{ --zemin:#f4f6f9; --kagit:#ffffff; --cizgi:#d5dce5; --baslik:#1f3a5f; --metin:#1c2430; --soluk:#5b6472; --secili:#2563eb; }}
* {{ box-sizing:border-box; }}
html, body {{ margin:0; height:100%; background:var(--zemin); color:var(--metin);
  font:14px/1.35 "DejaVu Sans", "Segoe UI", Arial, sans-serif; }}
header {{ display:flex; align-items:center; gap:14px; padding:10px 16px; background:var(--baslik); color:#fff; }}
header .marka {{ font-weight:700; letter-spacing:.2px; }}
header .dosya {{ font-family:"DejaVu Sans Mono", monospace; background:rgba(255,255,255,.12); padding:3px 8px; border-radius:6px; }}
header .kip {{ margin-left:auto; font-size:12px; opacity:.9; }}
header a {{ color:#fff; }}
.cubuk {{ display:flex; align-items:stretch; gap:0; background:var(--kagit); border-bottom:1px solid var(--cizgi); font-family:"DejaVu Sans Mono", monospace; }}
.cubuk .ad {{ width:84px; padding:7px 10px; border-right:1px solid var(--cizgi); font-weight:700; text-align:center; }}
.cubuk .fxi {{ padding:7px 10px; color:var(--soluk); border-right:1px solid var(--cizgi); font-style:italic; }}
.cubuk .icerik {{ padding:7px 10px; flex:1; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }}
.cubuk .tr {{ padding:7px 10px; color:var(--soluk); font-size:12px; white-space:nowrap; border-left:1px solid var(--cizgi); }}
main {{ height:calc(100% - 44px - 34px - 36px); overflow:auto; background:var(--kagit); }}
table {{ border-collapse:collapse; table-layout:fixed; }}
th, td {{ border:1px solid #e3e8ef; padding:0 5px; height:22px; overflow:hidden; white-space:nowrap; font-size:13px; }}
tr.bas th {{ background:#eef2f7; color:var(--soluk); font-weight:400; font-size:12px; position:sticky; top:0; z-index:2; }}
th.no {{ background:#eef2f7; color:var(--soluk); font-weight:400; font-size:12px; text-align:center; position:sticky; left:0; z-index:1; }}
td.sag {{ text-align:right; }} td.orta {{ text-align:center; }}
td.sar {{ white-space:normal; line-height:1.15; }}
td.eksi span {{ color:#b42318; }}
td.tasar {{ overflow:visible; }} td.tasar span {{ position:relative; z-index:1; }}
td.secili {{ outline:2px solid var(--secili); outline-offset:-2px; }}
body.formul td.fx span {{ font-family:"DejaVu Sans Mono", monospace; font-size:12px; color:#0b5394; }}
body.formul td.fx {{ text-align:left; }}
nav.sekme {{ display:flex; gap:2px; padding:0 8px; background:#e9edf2; border-top:1px solid var(--cizgi); height:36px; align-items:flex-end; overflow:hidden; }}
nav.sekme a {{ padding:7px 14px; background:#dfe5ec; color:var(--metin); text-decoration:none; border-radius:6px 6px 0 0; font-size:13px; }}
nav.sekme a.acik {{ background:var(--kagit); font-weight:700; box-shadow:inset 0 3px 0 var(--baslik); }}
section.sayfa {{ display:none; }} section.sayfa.acik {{ display:block; }}
</style></head>
<body>
<header><span class="marka">Finans 101 · Lab görüntüleyici</span><span class="dosya">calisma-kitaplari/{kitap}</span>
<span class="kip" id="kip">değerler · formül motoru: kitaplar/hesapla.py</span></header>
<div class="cubuk"><div class="ad" id="ad">A1</div><div class="fxi">fx</div><div class="icerik" id="icerik"></div><div class="tr" id="tr"></div></div>
<main id="ana">{sayfalar}</main>
<nav class="sekme">{sekmeler}</nav>
<script>
const q = new URLSearchParams(location.search);
const i = Math.max(1, parseInt(q.get('sayfa') || '1', 10));
const formul = q.get('formul') === '1';
const olcek = parseFloat(q.get('olcek') || '1');

if (formul) {{ document.body.classList.add('formul'); document.getElementById('kip').textContent = 'formüller gösteriliyor'; }}
const sayfalar = [...document.querySelectorAll('section.sayfa')];
const s = sayfalar[i - 1] || sayfalar[0];
s.classList.add('acik');
if (olcek !== 1) s.querySelector('table').style.zoom = olcek;
document.querySelectorAll('nav.sekme a')[sayfalar.indexOf(s)].classList.add('acik');
const araliklar = (q.get('satir') || '').split(',').filter(Boolean).map(a => a.split('-').map(Number));
if (araliklar.length) {{
  s.querySelectorAll('tr[data-r]').forEach(tr => {{ const r = +tr.dataset.r; if (!araliklar.some(([a, b]) => r >= a && r <= (b || a))) tr.style.display = 'none'; }});
}}
s.querySelectorAll('td[data-f]').forEach(td => {{ if (formul) td.firstChild.textContent = td.dataset.f; }});
// uzun etiket metni boş komşu hücrelere taşsın
s.querySelectorAll('tr').forEach(tr => {{
  const tds = [...tr.querySelectorAll('td')];
  tds.forEach((td, k) => {{ if (td.dataset.v && !td.classList.contains('sag') && !td.classList.contains('sar') && tds[k + 1] && !tds[k + 1].dataset.v && !(formul && td.dataset.f)) td.classList.add('tasar'); }});
}});
function sec(td) {{
  s.querySelectorAll('td.secili').forEach(x => x.classList.remove('secili'));
  td.classList.add('secili');
  document.getElementById('ad').textContent = td.dataset.a;
  document.getElementById('icerik').textContent = td.dataset.f || td.dataset.v || '';
  document.getElementById('tr').textContent = td.dataset.tr ? 'Türkçe Excel: ' + td.dataset.tr : '';
}}
s.addEventListener('click', e => {{ const td = e.target.closest('td'); if (td) sec(td); }});
const h = q.get('hucre');
const hedef = h && s.querySelector(`td[data-a="${{h.toUpperCase()}}"]`);
sec(hedef || s.querySelector('td[data-a="A1"]'));
document.body.dataset.hazir = '1';
</script>
</body></html>"""

INDEX = """<!DOCTYPE html>
<html lang="tr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Finans 101 · Lab görüntüleyici</title>
<style>
body {{ margin:0; background:#f4f6f9; color:#1c2430; font:15px/1.5 "DejaVu Sans", "Segoe UI", Arial, sans-serif; }}
header {{ background:#1f3a5f; color:#fff; padding:18px 24px; }} header h1 {{ margin:0; font-size:20px; }}
main {{ max-width:1100px; margin:0 auto; padding:18px 24px; }}
.uyari {{ background:#fff7e6; border:1px solid #f0c36d; border-radius:8px; padding:10px 14px; margin-bottom:16px; }}
.kitap {{ background:#fff; border:1px solid #d5dce5; border-radius:10px; padding:12px 16px; margin-bottom:12px; }}
.kitap h2 {{ margin:0 0 6px; font-size:16px; font-family:"DejaVu Sans Mono", monospace; }}
.kitap a {{ display:inline-block; margin:3px 8px 3px 0; padding:3px 10px; background:#e8eef5; border-radius:6px; color:#1f3a5f; text-decoration:none; }}
.kitap .k {{ color:#5b6472; font-size:13px; }}
</style></head><body>
<header><h1>Finans 101 · Lab görüntüleyici</h1></header>
<main>
<p class="uyari"><b>ÖRNEK VERİ.</b> Kişiler ve işletmeler uydurmadır; oranlar örnek değerdir, bugünkü piyasa bilgisi değildir.
Hafta 4 senaryoları gerçek ürün ya da piyasa verisi değildir. Bu lab <b>yatırım tavsiyesi değildir</b>.</p>
<p>Her dosya <code>calisma-kitaplari/</code> içindeki gerçek <code>.xlsx</code> dosyasıdır; Excel ya da Google E-Tablolar ile de açılır.
Burada gördüğünüz değerler, dosyadaki formüllerin formül motoruyla hesaplanmış hâlidir. Bir hücreye tıklayınca formülü üstte görünür.</p>
{kitaplar}
<p class="k">Oluşturma: <code>uv run python goruntuleyici/olustur.py</code> · kontrol: <code>./kontrol.sh</code></p>
</main></body></html>"""


def main() -> None:
    SITE.mkdir(parents=True, exist_ok=True)
    liste = []
    for yol in sorted((LAB / "calisma-kitaplari").glob("*.xlsx")):
        hesap = json.loads((LAB / "hesap" / (yol.stem + ".json")).read_text(encoding="utf-8"))["sayfalar"]
        wb = load_workbook(yol)
        sayfalar = "".join(sayfa_html(ws, hesap[ws.title], i) for i, ws in enumerate(wb.worksheets))
        sekmeler = "".join(f'<a href="?sayfa={i + 1}">{html.escape(ws.title)}</a>' for i, ws in enumerate(wb.worksheets))
        (SITE / (yol.stem + ".html")).write_text(SABLON.format(kitap=yol.name, sayfalar=sayfalar, sekmeler=sekmeler), encoding="utf-8")
        formul = sum(1 for s in hesap.values() for h in s.values() if h["f"])
        baglantilar = "".join(f'<a href="{yol.stem}.html?sayfa={i + 1}">{html.escape(ws.title)}</a>' for i, ws in enumerate(wb.worksheets))
        liste.append(f'<div class="kitap"><h2>{yol.name}</h2>{baglantilar}<div class="k">{len(wb.worksheets)} sayfa · {formul} formül</div></div>')
        print(f"goruntuleyici/site/{yol.stem}.html")
    (SITE / "index.html").write_text(INDEX.format(kitaplar="".join(liste)), encoding="utf-8")
    print("goruntuleyici/site/index.html")


if __name__ == "__main__":
    main()
