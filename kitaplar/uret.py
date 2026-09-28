"""Finans 101 lab çalışma kitaplarını (.xlsx) GERÇEK formüllerle üretir.

Kullanım (lab klasöründen):
    uv run python kitaplar/uret.py            → calisma-kitaplari/*.xlsx

- Sarı hücre = girdi (öğrenci değiştirir), beyaz hücre = formül. Formüller İngilizce işlev adlarıyla
  kaydedilir (dosya biçimi böyledir); Türkçe Excel aynı dosyayı TOPLA, EĞER, DEVRESEL_ÖDEME … diye gösterir.
- Bütün kişiler ve işletmeler UYDURMADIR (Deniz Örnek, Örnek Mum Atölyesi). Bütün oranlar ÖRNEK değerdir,
  bugünkü piyasa bilgisi değildir. Hafta 4 senaryoları uydurmadır; ürün, fon, hisse değildir.
- Değerler bu dosyada hesaplanmaz: hesapla.py formülleri çalıştırır, kontrol_degerler.py bağımsız olarak
  yeniden hesaplayıp karşılaştırır.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LAB = Path(__file__).resolve().parent.parent
CIKTI = LAB / "calisma-kitaplari"

TL0 = '#,##0 "TL"'
TL2 = '#,##0.00 "TL"'
YUZDE1 = "0.0%"
YUZDE0 = "0%"
ADET = "#,##0"
SAYI = "0"

SARI = PatternFill("solid", fgColor="FFF3C4")
BASLIK = PatternFill("solid", fgColor="1F3A5F")
ARA = PatternFill("solid", fgColor="E8EEF5")
INCE = Side(style="thin", color="C9D3DF")
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)

AYLAR = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
AY_SUTUN = [get_column_letter(i) for i in range(2, 14)]  # B..M

NOT_KISI = "ÖRNEK VERİ · Deniz Örnek uydurma bir kişidir · oranlar ÖRNEK değerdir, bugünkü piyasa bilgisi değildir"
NOT_ISLETME = "ÖRNEK VERİ · Örnek Mum Atölyesi uydurma bir işletmedir · tutarlar ve oranlar ÖRNEKTİR · vergiler bu tabloda yoktur"
NOT_YATIRIM = ("UYDURMA SENARYO · gerçek ürün, fon, hisse ya da geçmiş piyasa verisi DEĞİLDİR · "
               "bu tablo YATIRIM TAVSİYESİ DEĞİLDİR")
NOT_RENK = "Sarı hücre = girdi (değiştirebilirsiniz) · beyaz hücre = formül (elle yazmayın)"


class Sayfa:
    def __init__(self, wb: Workbook, ad: str, baslik: str, not_: str, genislik: dict[str, float]):
        self.ws = wb.create_sheet(ad)
        ws = self.ws
        ws["A1"] = baslik
        ws["A1"].font = Font(bold=True, size=14, color="1F3A5F")
        ws["A2"] = not_
        ws["A2"].font = Font(italic=True, size=9, color="9A3412")
        ws["A3"] = NOT_RENK
        ws["A3"].font = Font(size=9, color="5B6472")
        for sut, g in genislik.items():
            ws.column_dimensions[sut].width = g
        ws.sheet_view.showGridLines = True

    def etiket(self, adres: str, metin: str, kalin: bool = False):
        c = self.ws[adres]
        c.value = metin
        if kalin:
            c.font = Font(bold=True)
        return c

    def girdi(self, adres: str, deger, bicim: str | None = None):
        c = self.ws[adres]
        c.value = deger
        c.fill = SARI
        c.border = KENAR
        if bicim:
            c.number_format = bicim
        return c

    def formul(self, adres: str, f: str, bicim: str | None = None, kalin: bool = False):
        assert f.startswith("="), adres
        c = self.ws[adres]
        c.value = f
        c.border = KENAR
        if bicim:
            c.number_format = bicim
        if kalin:
            c.font = Font(bold=True)
        return c

    def baslik_satiri(self, satir: int, basliklar: list[str], ilk_sutun: int = 1):
        for i, b in enumerate(basliklar):
            c = self.ws.cell(row=satir, column=ilk_sutun + i, value=b)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = BASLIK
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = KENAR
        self.ws.row_dimensions[satir].height = 30

    def ara(self, adres: str, metin: str):
        c = self.ws[adres]
        c.value = metin
        c.font = Font(bold=True)
        c.fill = ARA
        c.border = KENAR
        return c

    def not_(self, adres: str, metin: str):
        c = self.ws[adres]
        c.value = metin
        c.font = Font(italic=True, size=9, color="5B6472")
        return c


def yeni_kitap() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True  # Excel/Sheets açılışta her formülü yeniden hesaplar
    return wb


# ============================================================ HAFTA 1 · Paranın Temelleri
ISLEMLER = [  # gün, açıklama, kategori, tür, tutar — Deniz Örnek'in örnek ayı
    (1, "Maaş", "Maaş", "Gelir", 42000),
    (1, "Kira", "Kira", "Gider", 15000),
    (2, "Market alışverişi", "Market", "Gider", 2100),
    (3, "Elektrik, su, doğalgaz", "Faturalar", "Gider", 2800),
    (4, "Toplu taşıma kartı", "Ulaşım", "Gider", 1200),
    (5, "Telefon ve internet", "Faturalar", "Gider", 900),
    (7, "Kafe", "Dışarıda yemek", "Gider", 450),
    (9, "Market alışverişi", "Market", "Gider", 2600),
    (10, "Serbest iş ödemesi", "Ek gelir", "Gelir", 3500),
    (12, "Dizi aboneliği", "Abonelik", "Gider", 250),
    (14, "Giyim", "Giyim", "Gider", 1800),
    (16, "Market alışverişi", "Market", "Gider", 2300),
    (18, "Akşam yemeği", "Dışarıda yemek", "Gider", 1400),
    (20, "Taksi", "Ulaşım", "Gider", 650),
    (21, "Spor salonu", "Abonelik", "Gider", 1100),
    (22, "Kitap", "Diğer", "Gider", 350),
    (23, "Market alışverişi", "Market", "Gider", 1900),
    (25, "Hediye", "Diğer", "Gider", 1200),
    (27, "Eczane", "Sağlık", "Gider", 600),
    (29, "Birikim hesabına aktarım", "Birikim", "Birikim", 4000),
]
KATEGORILER = ["Kira", "Faturalar", "Market", "Ulaşım", "Dışarıda yemek", "Abonelik", "Giyim", "Sağlık", "Diğer"]
ZORUNLU = {"Kira": "Evet", "Faturalar": "Evet", "Market": "Evet", "Ulaşım": "Evet", "Dışarıda yemek": "Hayır",
           "Abonelik": "Hayır", "Giyim": "Hayır", "Sağlık": "Evet", "Diğer": "Hayır"}
PLAN = {"Kira": 15000, "Faturalar": 3500, "Market": 8000, "Ulaşım": 2000, "Dışarıda yemek": 1500,
        "Abonelik": 1200, "Giyim": 1000, "Sağlık": 700, "Diğer": 800}


def hafta1() -> Workbook:
    wb = yeni_kitap()

    # --- 1.1 Gelir-Gider
    s = Sayfa(wb, "Gelir-Gider", "Gelir-Gider Kaydı · bir ay (1.1)", NOT_KISI,
              {"A": 7, "B": 26, "C": 16, "D": 10, "E": 14, "F": 3, "G": 30, "H": 14})
    s.baslik_satiri(5, ["Gün", "Açıklama", "Kategori", "Tür", "Tutar"])
    for i, (gun, aciklama, kat, tur, tutar) in enumerate(ISLEMLER):
        r = 6 + i
        s.girdi(f"A{r}", gun, SAYI)
        s.girdi(f"B{r}", aciklama)
        s.girdi(f"C{r}", kat)
        s.girdi(f"D{r}", tur)
        s.girdi(f"E{r}", tutar, TL0)
    s.baslik_satiri(5, ["Özet", "Tutar"], ilk_sutun=7)
    s.etiket("G6", "Toplam gelir")
    s.formul("H6", '=SUMIF($D$6:$D$25,"Gelir",$E$6:$E$25)', TL0)
    s.etiket("G7", "Toplam gider")
    s.formul("H7", '=SUMIF($D$6:$D$25,"Gider",$E$6:$E$25)', TL0)
    s.etiket("G8", "Birikime ayrılan")
    s.formul("H8", '=SUMIF($D$6:$D$25,"Birikim",$E$6:$E$25)', TL0)
    s.ara("G9", "Kalan (gelir − gider − birikim)")
    s.formul("H9", "=H6-H7-H8", TL0, kalin=True)
    s.etiket("G10", "Gelirin birikime giden payı")
    s.formul("H10", "=IFERROR(H8/H6,0)", YUZDE1)
    s.baslik_satiri(12, ["Kategori", "Bu ay"], ilk_sutun=7)
    for i, k in enumerate(KATEGORILER):
        r = 13 + i
        s.etiket(f"G{r}", k)
        s.formul(f"H{r}", f"=SUMIF($C$6:$C$25,G{r},$E$6:$E$25)", TL0)
    s.ara("G22", "Kategoriler toplamı")
    s.formul("H22", "=SUM(H13:H21)", TL0, kalin=True)
    s.etiket("G23", "Toplam giderle aynı mı?")
    s.formul("H23", '=IF(ROUND(H22-H7,2)=0,"EVET","HAYIR — kategori eksik")')
    s.ws.freeze_panes = "A6"

    # --- 1.2 Bütçe
    s = Sayfa(wb, "Bütçe", "Aylık Bütçe · plan ve gerçekleşen (1.2)", NOT_KISI,
              {"A": 34, "B": 12, "C": 14, "D": 16, "E": 18, "F": 13, "G": 12})
    s.baslik_satiri(5, ["Kategori", "Zorunlu mu?", "Plan", "Gerçekleşen", "Fark (plan − gerçek)", "Gelire oranı", "Durum"])
    for i, k in enumerate(KATEGORILER):
        r = 6 + i
        s.etiket(f"A{r}", k)
        s.girdi(f"B{r}", ZORUNLU[k])
        s.girdi(f"C{r}", PLAN[k], TL0)
        s.formul(f"D{r}", f"=SUMIF('Gelir-Gider'!$C$6:$C$25,A{r},'Gelir-Gider'!$E$6:$E$25)", TL0)
        s.formul(f"E{r}", f"=C{r}-D{r}", TL0)
        s.formul(f"F{r}", f"=IFERROR(D{r}/$C$19,0)", YUZDE1)
        s.formul(f"G{r}", f'=IF(E{r}<0,"AŞILDI","")')
    s.ara("A15", "Toplam gider")
    s.formul("C15", "=SUM(C6:C14)", TL0, kalin=True)
    s.formul("D15", "=SUM(D6:D14)", TL0, kalin=True)
    s.formul("E15", "=C15-D15", TL0, kalin=True)
    s.formul("F15", "=IFERROR(D15/$C$19,0)", YUZDE1)
    s.etiket("A17", "Zorunlu giderler toplamı")
    s.formul("C17", '=SUMIF($B$6:$B$14,"Evet",C6:C14)', TL0)
    s.formul("D17", '=SUMIF($B$6:$B$14,"Evet",D6:D14)', TL0)
    s.etiket("A19", "Aylık gelir (Gelir-Gider sayfasından)")
    s.formul("C19", "='Gelir-Gider'!H6", TL0)
    s.formul("D19", "='Gelir-Gider'!H6", TL0)
    s.etiket("A20", "Birikime ayrılan")
    s.girdi("C20", 4000, TL0)
    s.formul("D20", "='Gelir-Gider'!H8", TL0)
    s.ara("A21", "Kalan (gelir − gider − birikim)")
    s.formul("C21", "=C19-C15-C20", TL0, kalin=True)
    s.formul("D21", "=D19-D15-D20", TL0, kalin=True)
    s.not_("A23", "Plan sütununu siz yazarsınız; gerçekleşen sütunu Gelir-Gider sayfasındaki kayıtlardan kendiliğinden gelir.")

    # --- 1.3 Nakit Akışı (12 ay)
    s = Sayfa(wb, "Nakit Akışı", "12 Aylık Kişisel Nakit Akışı · örnek yıl (1.3)", NOT_KISI,
              {"A": 32, **{c: 10.5 for c in AY_SUTUN}, "N": 12.5})
    s.baslik_satiri(5, ["Kalem", *AYLAR, "Yıl toplamı"])
    duzensiz = {"M": 15000, "L": 5000}           # Aralık ikramiye, Kasım ek iş (örnek)
    donemsel = {"D": 14000, "H": 25000, "J": 12000, "M": 6000}  # Mart sigorta, Temmuz tatil, Eylül kurs, Aralık yılbaşı
    etiketler = {6: "Açılış bakiyesi", 7: "Düzenli gelir (Bütçe)", 8: "Düzensiz gelir", 9: "Düzenli gider (Bütçe)",
                 10: "Dönemsel gider (sigorta, tatil …)", 11: "Birikime ayrılan (Bütçe)", 12: "Kapanış bakiyesi", 13: "Durum"}
    for r, e in etiketler.items():
        (s.ara if r == 12 else s.etiket)(f"A{r}", e)
    for i, c in enumerate(AY_SUTUN):
        onceki = AY_SUTUN[i - 1] if i else None
        if onceki:
            s.formul(f"{c}6", f"={onceki}12", TL0)
        else:
            s.girdi(f"{c}6", 2000, TL0)
        s.formul(f"{c}7", "='Bütçe'!$C$19", TL0)
        s.girdi(f"{c}8", duzensiz.get(c, 0), TL0)
        s.formul(f"{c}9", "='Bütçe'!$D$15", TL0)
        s.girdi(f"{c}10", donemsel.get(c, 0), TL0)
        s.formul(f"{c}11", "='Bütçe'!$C$20", TL0)
        s.formul(f"{c}12", f"={c}6+{c}7+{c}8-{c}9-{c}10-{c}11", TL0, kalin=True)
        s.formul(f"{c}13", f'=IF({c}12<0,"AÇIK","")')
    for r in (7, 8, 9, 10, 11):
        s.formul(f"N{r}", f"=SUM(B{r}:M{r})", TL0)
    s.etiket("A15", "En düşük kapanış")
    s.formul("B15", "=MIN(B12:M12)", TL0)
    s.etiket("A16", "Açık veren ay sayısı")
    s.formul("B16", '=COUNTIF(B12:M12,"<0")', SAYI)
    s.not_("A18", "Düzensiz gelir ve dönemsel gider satırları sarı: yıl içinde ne zaman geleceğini siz yazarsınız.")
    s.ws.freeze_panes = "B6"

    # --- 1.4 Acil Durum Fonu
    s = Sayfa(wb, "Acil Durum Fonu", "Acil Durum Fonu Hedefi (1.4)", NOT_KISI,
              {"A": 38, "B": 13, "C": 13, "D": 13, "E": 22})
    s.etiket("A5", "Zorunlu aylık gider (Bütçe sayfasından)")
    s.formul("C5", "='Bütçe'!D17", TL0)
    s.etiket("A6", "Hedef: kaç aylık zorunlu gider?")
    s.girdi("C6", 3, SAYI)
    s.not_("D6", "Kendi durumunuza göre seçin — kurs sabit bir sayı önermez.")
    s.ara("A7", "Hedef tutar")
    s.formul("C7", "=C5*C6", TL0, kalin=True)
    s.etiket("A8", "Şu an ayrılmış tutar")
    s.girdi("C8", 12000, TL0)
    s.etiket("A9", "Eksik")
    s.formul("C9", "=MAX(0,C7-C8)", TL0)
    s.etiket("A10", "Her ay ayırabileceğim")
    s.girdi("C10", 4000, TL0)
    s.ara("A11", "Hedefe kaç ayda ulaşırım?")
    s.formul("C11", "=IF(C9=0,0,ROUNDUP(C9/C10,0))", SAYI, kalin=True)
    s.baslik_satiri(13, ["Ay", "Ay başı", "Eklenen", "Ay sonu", "Hedefe ulaşıldı mı?"])
    for i in range(24):
        r = 14 + i
        s.ws[f"A{r}"] = i + 1
        s.formul(f"B{r}", "=C8" if i == 0 else f"=D{r - 1}", TL0)
        s.formul(f"C{r}", f"=MIN($C$10,MAX(0,$C$7-B{r}))", TL0)
        s.formul(f"D{r}", f"=B{r}+C{r}", TL0)
        s.formul(f"E{r}", f'=IF(D{r}>=$C$7,"EVET","")')
    return wb


# ============================================================ HAFTA 2 · Borç ve Birikim
def hafta2() -> Workbook:
    wb = yeni_kitap()

    # --- 2.1 Bileşik Faiz
    s = Sayfa(wb, "Bileşik Faiz", "Basit ve Bileşik Faiz (2.1)", NOT_KISI,
              {"A": 34, "B": 16, "C": 16, "D": 16})
    s.etiket("A5", "Anapara")
    s.girdi("C5", 100000, TL0)
    s.etiket("A6", "Yıllık faiz oranı (ÖRNEK)")
    s.girdi("C6", 0.25, YUZDE1)
    s.etiket("A7", "Süre (yıl)")
    s.girdi("C7", 10, SAYI)
    s.baslik_satiri(9, ["Yıl", "Basit faizle", "Bileşik faizle", "Fark"])
    for n in range(11):
        r = 10 + n
        s.ws[f"A{r}"] = n
        s.formul(f"B{r}", f"=$C$5*(1+$C$6*A{r})", TL0)
        s.formul(f"C{r}", f"=$C$5*(1+$C$6)^A{r}", TL0)
        s.formul(f"D{r}", f"=C{r}-B{r}", TL0)
    s.etiket("A22", "Kontrol: FV (GD) işleviyle son yıl")
    s.formul("C22", "=FV(C6,C7,0,-C5)", TL0)
    s.not_("A24", "Basit faiz yalnız anaparaya işler; bileşik faizde her yılın faizi bir sonraki yıl anaparaya eklenir.")

    # --- 2.2a Kredi Kartı
    s = Sayfa(wb, "Kredi Kartı", "Kredi Kartı: Bu Ay Ne Kadar Ödemeli? (2.2)", NOT_KISI,
              {"A": 30, "B": 13, "C": 13, "D": 14, "E": 12, "F": 14, "G": 14})
    s.etiket("A5", "Dönem borcu (hesap özeti)")
    s.girdi("C5", 20000, TL0)
    s.etiket("A6", "Aylık akdi faiz (ÖRNEK)")
    s.girdi("C6", 0.04, YUZDE1)
    s.etiket("A7", "Aylık gecikme faizi (ÖRNEK)")
    s.girdi("C7", 0.043, YUZDE1)
    s.etiket("A8", "Asgari ödeme oranı (sözleşmenizdeki)")
    s.girdi("C8", 0.2, YUZDE0)
    s.baslik_satiri(10, ["Seçenek", "Ödeme oranı", "Ödenen", "Kalan bakiye", "Faiz türü", "Aylık oran", "Bu ayın faizi"])
    secenekler = [("Tamamını öde", 1), ("Yarısını öde", 0.5), ("Asgari tutarı öde", "=C8"), ("Asgari tutarın altında öde", 0.1)]
    for i, (ad, oran) in enumerate(secenekler):
        r = 11 + i
        s.etiket(f"A{r}", ad)
        if isinstance(oran, str):
            s.formul(f"B{r}", oran, YUZDE0)
        else:
            s.girdi(f"B{r}", oran, YUZDE0)
        s.formul(f"C{r}", f"=$C$5*B{r}", TL0)
        s.formul(f"D{r}", f"=$C$5-C{r}", TL0)
        s.formul(f"E{r}", f'=IF(D{r}=0,"—",IF(B{r}>=$C$8,"akdi","gecikme"))')
        s.formul(f"F{r}", f"=IF(D{r}=0,0,IF(B{r}>=$C$8,$C$6,$C$7))", YUZDE1)
        s.formul(f"G{r}", f"=D{r}*F{r}", TL0, kalin=True)
    s.not_("A16", "Kavram modeli (YAKLAŞIK): tek hesap dönemi; vergi, ücret ve bankanın gün hesabı yok. Kalan bakiyeye faiz işler;")
    s.not_("A17", "asgari tutarın altında ödemede gecikme faizi uygulanır (5464 sayılı Kanun md. 26). Gerçek tutar hesap özetinizdedir.")

    # --- 2.2b Kredi Planı
    s = Sayfa(wb, "Kredi Planı", "Kredi Geri Ödeme Planı (2.2)", NOT_KISI,
              {"A": 30, "B": 15, "C": 13, "D": 13, "E": 13, "F": 15})
    s.etiket("A5", "Kredi tutarı")
    s.girdi("C5", 50000, TL0)
    s.etiket("A6", "Aylık faiz (ÖRNEK)")
    s.girdi("C6", 0.03, YUZDE1)
    s.etiket("A7", "Vade (ay) — tablo 12 ay içindir")
    s.girdi("C7", 12, SAYI)
    s.ara("A8", "Aylık taksit")
    s.formul("C8", "=-PMT(C6,C7,C5)", TL2, kalin=True)
    s.etiket("A9", "Toplam geri ödeme")
    s.formul("C9", "=C8*C7", TL2)
    s.ara("A10", "Toplam faiz (kredinin faiz maliyeti)")
    s.formul("C10", "=C9-C5", TL2, kalin=True)
    s.baslik_satiri(12, ["Ay", "Ay başı borç", "Taksit", "Faiz", "Anapara", "Ay sonu borç"])
    for i in range(12):
        r = 13 + i
        s.ws[f"A{r}"] = i + 1
        s.formul(f"B{r}", "=C5" if i == 0 else f"=F{r - 1}", TL2)
        s.formul(f"C{r}", "=$C$8", TL2)
        s.formul(f"D{r}", f"=B{r}*$C$6", TL2)
        s.formul(f"E{r}", f"=C{r}-D{r}", TL2)
        s.formul(f"F{r}", f"=B{r}-E{r}", TL2)
    s.ara("A25", "Toplam")
    for c in "CDE":
        s.formul(f"{c}25", f"=SUM({c}13:{c}24)", TL2, kalin=True)
    s.etiket("A27", "Kontrol: IPMT (FAİZTUTARI) ile 1. ay faizi")
    s.formul("C27", "=-IPMT(C6,1,C7,C5)", TL2)
    s.not_("A29", "Vergi, masraf ve sigorta bu tabloda yok. Kredinin toplam maliyeti sözleşme öncesi bilgi formunda yazar (6502 sayılı Kanun md. 23).")

    # --- 2.3 Birikim Hedefi
    s = Sayfa(wb, "Birikim Hedefi", "Birikim Hedefi: Ayda Ne Kadar? (2.3)", NOT_KISI,
              {"A": 34, "B": 14, "C": 14, "D": 13, "E": 15})
    s.etiket("A5", "Hedef tutar")
    s.girdi("C5", 120000, TL0)
    s.etiket("A6", "Süre (ay) — tablo 24 ay içindir")
    s.girdi("C6", 24, SAYI)
    s.etiket("A7", "Birikimin aylık getirisi (ÖRNEK)")
    s.girdi("C7", 0.015, YUZDE1)
    s.etiket("A8", "Şu anki birikim")
    s.girdi("C8", 10000, TL0)
    s.ara("A9", "Gereken aylık katkı")
    s.formul("C9", "=-PMT(C7,C6,-C8,C5)", TL2, kalin=True)
    s.etiket("A10", "Getiri olmasaydı gereken aylık katkı")
    s.formul("C10", "=(C5-C8)/C6", TL2)
    s.etiket("A11", "Getirinin toplam katkısı")
    s.formul("C11", "=C5-C8-C9*C6", TL2)
    s.baslik_satiri(13, ["Ay", "Ay başı", "Katkı", "Getiri", "Ay sonu"])
    for i in range(24):
        r = 14 + i
        s.ws[f"A{r}"] = i + 1
        s.formul(f"B{r}", "=C8" if i == 0 else f"=E{r - 1}", TL2)
        s.formul(f"C{r}", "=$C$9", TL2)
        s.formul(f"D{r}", f"=B{r}*$C$7", TL2)
        s.formul(f"E{r}", f"=B{r}+C{r}+D{r}", TL2)

    # --- 2.4 Enflasyon
    s = Sayfa(wb, "Enflasyon", "Enflasyonun Etkisi: Satın Alma Gücü (2.4)", NOT_KISI,
              {"A": 44, "B": 15, "C": 18, "D": 18, "E": 20})
    s.etiket("A5", "Bugünkü sepet tutarı")
    s.girdi("C5", 10000, TL0)
    s.etiket("A6", "Yıllık enflasyon (ÖRNEK)")
    s.girdi("C6", 0.2, YUZDE1)
    s.etiket("A7", "Birikimin yıllık faizi (ÖRNEK)")
    s.girdi("C7", 0.22, YUZDE1)
    s.baslik_satiri(9, ["Yıl", "Aynı sepetin fiyatı", "Paranın alım gücü (bugünkü TL)", "Birikimin tutarı", "Birikimin alım gücü (bugünkü TL)"])
    for n in range(6):
        r = 10 + n
        s.ws[f"A{r}"] = n
        s.formul(f"B{r}", f"=$C$5*(1+$C$6)^A{r}", TL0)
        s.formul(f"C{r}", f"=$C$5/(1+$C$6)^A{r}", TL0)
        s.formul(f"D{r}", f"=$C$5*(1+$C$7)^A{r}", TL0)
        s.formul(f"E{r}", f"=D{r}/(1+$C$6)^A{r}", TL0)
    s.ara("A17", "Reel getiri (yıllık, enflasyondan arındırılmış)")
    s.formul("C17", "=(1+C7)/(1+C6)-1", YUZDE1, kalin=True)
    s.etiket("A18", "Birikim Hedefi 2 yıl sonra kaç TL eder?")
    s.formul("C18", "='Birikim Hedefi'!C5*(1+C6)^2", TL0)
    s.etiket("A19", "Bu tutar için gereken aylık katkı")
    s.formul("C19", "=-PMT('Birikim Hedefi'!C7,'Birikim Hedefi'!C6,-'Birikim Hedefi'!C8,C18)", TL2)
    s.not_("A21", "Resmî enflasyon verisi TÜİK'in Tüketici Fiyat Endeksi'dir (TÜFE); buradaki oran yalnız örnektir.")
    return wb


# ============================================================ HAFTA 3 · İşletme Finansı
def hafta3() -> Workbook:
    wb = yeni_kitap()
    alti = AY_SUTUN[:6]  # B..G

    # --- 3.1 Kâr ve Nakit
    s = Sayfa(wb, "Kâr ve Nakit", "Kâr ile Nakit Aynı Şey Değil · ilk 6 ay (3.1)", NOT_ISLETME,
              {"A": 36, **{c: 15.5 for c in alti}, "H": 13})
    s.etiket("A5", "Satışın peşin tahsil edilen oranı")
    s.girdi("C5", 0.4, YUZDE0)
    s.not_("D5", "Kalanı iki ay sonra tahsil edilir (vadeli satış).")
    s.baslik_satiri(8, ["Kalem", *AYLAR[:6], "Toplam"])
    satis = [40000, 55000, 70000, 80000, 90000, 100000]
    gider = [45000, 50000, 55000, 60000, 62000, 65000]
    et = {9: "Satış (kesilen fatura)", 10: "Giderler (aynı ay ödenen)", 11: "Kâr (satış − gider)",
          12: "Tahsilat: peşin kısım", 13: "Tahsilat: iki ay önceki satışın kalanı", 14: "Nakit girişi",
          15: "Nakit akışı (giriş − gider)", 16: "Birikimli kâr", 17: "Birikimli nakit", 18: "Kâr artıda, nakit ekside mi?"}
    for r, e in et.items():
        (s.ara if r in (11, 15) else s.etiket)(f"A{r}", e)
    for i, c in enumerate(alti):
        s.girdi(f"{c}9", satis[i], TL0)
        s.girdi(f"{c}10", gider[i], TL0)
        s.formul(f"{c}11", f"={c}9-{c}10", TL0, kalin=True)
        s.formul(f"{c}12", f"={c}9*$C$5", TL0)
        if i < 2:
            s.girdi(f"{c}13", 0, TL0)
        else:
            s.formul(f"{c}13", f"={alti[i - 2]}9*(1-$C$5)", TL0)
        s.formul(f"{c}14", f"={c}12+{c}13", TL0)
        s.formul(f"{c}15", f"={c}14-{c}10", TL0, kalin=True)
        s.formul(f"{c}16", f"={c}11" if i == 0 else f"={alti[i - 1]}16+{c}11", TL0)
        s.formul(f"{c}17", f"={c}15" if i == 0 else f"={alti[i - 1]}17+{c}15", TL0)
        s.formul(f"{c}18", f'=IF(AND({c}16>0,{c}17<0),"EVET","")')
    for r in (9, 10, 11, 12, 13, 14, 15):
        s.formul(f"H{r}", f"=SUM(B{r}:G{r})", TL0)
    s.ara("A20", "Başlangıçta gereken nakit (en düşük birikimli nakit)")
    s.formul("C20", "=-MIN(B17:G17)", TL0, kalin=True)
    s.not_("A22", "Kâr, faturanın kesildiği ayda; nakit, paranın kasaya girdiği ayda görünür.")

    # --- 3.2 Başabaş
    s = Sayfa(wb, "Başabaş", "Başabaş Noktası (3.2)", NOT_ISLETME,
              {"A": 32, "B": 14, "C": 16, "D": 14, "E": 15, "F": 15})
    s.etiket("A5", "Aylık sabit giderler (kira, maaş …)")
    s.girdi("C5", 30000, TL0)
    s.etiket("A6", "Birim satış fiyatı (bir mum)")
    s.girdi("C6", 250, TL0)
    s.etiket("A7", "Birim değişken maliyet")
    s.girdi("C7", 130, TL0)
    s.etiket("A8", "Birim katkı payı")
    s.formul("C8", "=C6-C7", TL0)
    s.etiket("A9", "Katkı payı oranı")
    s.formul("C9", "=C8/C6", YUZDE1)
    s.ara("A10", "Başabaş adedi (aylık)")
    s.formul("C10", "=ROUNDUP(C5/C8,0)", ADET, kalin=True)
    s.etiket("A11", "Başabaş cirosu")
    s.formul("C11", "=C10*C6", TL0)
    s.etiket("A12", "Beklenen satış (adet)")
    s.girdi("C12", 400, ADET)
    s.etiket("A13", "Güvenlik payı (adet)")
    s.formul("C13", "=C12-C10", ADET)
    s.ara("A14", "Beklenen aylık kâr")
    s.formul("C14", "=C12*C8-C5", TL0, kalin=True)
    s.baslik_satiri(16, ["Adet", "Ciro", "Değişken maliyet", "Sabit gider", "Toplam maliyet", "Kâr / zarar"])
    for i in range(11):
        r = 17 + i
        s.ws[f"A{r}"] = i * 50
        s.ws[f"A{r}"].number_format = ADET
        s.formul(f"B{r}", f"=A{r}*$C$6", TL0)
        s.formul(f"C{r}", f"=A{r}*$C$7", TL0)
        s.formul(f"D{r}", "=$C$5", TL0)
        s.formul(f"E{r}", f"=C{r}+D{r}", TL0)
        s.formul(f"F{r}", f"=B{r}-E{r}", TL0, kalin=True)

    # --- 3.3 Fiyatlama
    s = Sayfa(wb, "Fiyatlama", "Fiyatlama: Maliyet, Marj, Başabaş (3.3)", NOT_ISLETME,
              {"A": 40, "B": 14, "C": 15, "D": 18, "E": 12})
    s.etiket("A5", "Birim değişken maliyet (Başabaş sayfasından)")
    s.formul("C5", "='Başabaş'!C7", TL0)
    s.etiket("A6", "Aylık sabit gider (Başabaş sayfasından)")
    s.formul("C6", "='Başabaş'!C5", TL0)
    s.etiket("A7", "Beklenen satış (adet)")
    s.formul("C7", "='Başabaş'!C12", ADET)
    s.etiket("A8", "Birim başına sabit gider payı")
    s.formul("C8", "=C6/C7", TL2)
    s.ara("A9", "Birim tam maliyet")
    s.formul("C9", "=C5+C8", TL2, kalin=True)
    s.etiket("A10", "Maliyetin üstüne eklenen oran (ÖRNEK)")
    s.girdi("C10", 0.4, YUZDE0)
    s.ara("A11", "Maliyet artı fiyat")
    s.formul("C11", "=C9*(1+C10)", TL2, kalin=True)
    s.etiket("A12", "Bu fiyatta kâr marjı (fiyata göre)")
    s.formul("C12", "=(C11-C9)/C11", YUZDE1)
    s.baslik_satiri(14, ["Fiyat seçeneği", "Katkı payı", "Başabaş adedi", "Beklenen adette kâr", "Kâr marjı"])
    for i, fiyat in enumerate([200, 220, 250, "=C11", 320]):
        r = 15 + i
        if isinstance(fiyat, str):
            s.formul(f"A{r}", fiyat, TL2)
        else:
            s.girdi(f"A{r}", fiyat, TL0)
        s.formul(f"B{r}", f"=A{r}-$C$5", TL2)
        s.formul(f"C{r}", f'=IFERROR(ROUNDUP($C$6/B{r},0),"—")', ADET)
        s.formul(f"D{r}", f"=$C$7*B{r}-$C$6", TL0)
        s.formul(f"E{r}", f"=IFERROR(D{r}/($C$7*A{r}),0)", YUZDE1)
    s.not_("A21", "Beklenen adet her fiyatta aynı varsayıldı; gerçekte fiyat değişince talep de değişir. KDV vb. vergiler tabloda yok.")

    # --- 3.4 Nakit Akış Tablosu
    s = Sayfa(wb, "Nakit Akış Tablosu", "12 Aylık İşletme Nakit Akış Tablosu (3.4)", NOT_ISLETME,
              {"A": 34, **{c: 10.5 for c in AY_SUTUN}, "N": 12})
    s.etiket("A5", "Birim fiyat (Başabaş sayfasından)")
    s.formul("C5", "='Başabaş'!C6", TL0)
    s.etiket("A6", "Birim değişken maliyet (Başabaş sayfasından)")
    s.formul("C6", "='Başabaş'!C7", TL0)
    s.etiket("A7", "Satışın peşin tahsil edilen oranı (kalanı ertesi ay)")
    s.girdi("C7", 0.5, YUZDE0)
    s.baslik_satiri(9, ["Kalem", *AYLAR, "Yıl toplamı"])
    adet = [250, 260, 300, 320, 350, 300, 280, 300, 380, 420, 480, 600]
    yatirim = {"E": 40000}  # Nisan: ekipman (örnek)
    et = {10: "Satış adedi", 11: "Satış tutarı", 12: "Tahsilat: bu ayın peşin kısmı", 13: "Tahsilat: geçen ayın kalanı",
          14: "Nakit girişi", 15: "Hammadde ödemesi", 16: "Sabit giderler", 17: "Ekipman / yatırım",
          18: "Nakit çıkışı", 19: "Net nakit akışı", 20: "Açılış bakiyesi", 21: "Kapanış bakiyesi", 22: "Durum"}
    for r, e in et.items():
        (s.ara if r in (19, 21) else s.etiket)(f"A{r}", e)
    for i, c in enumerate(AY_SUTUN):
        o = AY_SUTUN[i - 1] if i else None
        s.girdi(f"{c}10", adet[i], ADET)
        s.formul(f"{c}11", f"={c}10*$C$5", TL0)
        s.formul(f"{c}12", f"={c}11*$C$7", TL0)
        if o:
            s.formul(f"{c}13", f"={o}11*(1-$C$7)", TL0)
        else:
            s.girdi(f"{c}13", 0, TL0)
        s.formul(f"{c}14", f"={c}12+{c}13", TL0)
        s.formul(f"{c}15", f"={c}10*$C$6", TL0)
        s.formul(f"{c}16", "='Başabaş'!$C$5", TL0)
        s.girdi(f"{c}17", yatirim.get(c, 0), TL0)
        s.formul(f"{c}18", f"={c}15+{c}16+{c}17", TL0)
        s.formul(f"{c}19", f"={c}14-{c}18", TL0, kalin=True)
        if o:
            s.formul(f"{c}20", f"={o}21", TL0)
        else:
            s.girdi(f"{c}20", 50000, TL0)
        s.formul(f"{c}21", f"={c}20+{c}19", TL0, kalin=True)
        s.formul(f"{c}22", f'=IF({c}21<0,"KASA AÇIĞI","")')
    for r in (10, 11, 14, 15, 16, 17, 18, 19):
        s.formul(f"N{r}", f"=SUM(B{r}:M{r})", ADET if r == 10 else TL0)
    s.etiket("A24", "En düşük kapanış")
    s.formul("C24", "=MIN(B21:M21)", TL0)
    s.etiket("A25", "Kasa açığı veren ay sayısı")
    s.formul("C25", '=COUNTIF(B21:M21,"<0")', SAYI)
    s.ws.freeze_panes = "B10"
    return wb


# ============================================================ HAFTA 4 · Karar Vermek
SENARYO = {
    "A": [0.21, 0.22, 0.20, 0.23, 0.21, 0.22, 0.20, 0.21, 0.22, 0.21],
    "B": [0.30, 0.12, 0.35, -0.05, 0.28, 0.18, 0.40, 0.02, 0.25, 0.22],
    "C": [0.80, -0.35, 0.95, -0.40, 0.60, 0.10, 1.10, -0.50, 0.70, 0.05],
}
X = [0.30, -0.10, 0.25, 0.05, 0.35, -0.15, 0.28, 0.02, 0.30, -0.05]
Y = [-0.02, 0.25, 0.00, 0.20, -0.05, 0.30, 0.02, 0.18, 0.00, 0.22]
HEDEFLER = [  # ad, öncelik, bugünkü tutar, ay, plana dahil
    ("Acil durum fonu (kalan kısım)", 1, 60000, 12, "Evet"),
    ("Eğitim / sertifika", 2, 30000, 12, "Evet"),
    ("Dizüstü bilgisayar", 3, 45000, 18, "Hayır"),
    ("Tatil", 4, 25000, 10, "Hayır"),
]


def hedef_sayfasi(wb: Workbook, ad: str, baslik: str, not_: str, hedefler, enf, getiri, ayirabilen):
    s = Sayfa(wb, ad, baslik, not_, {"A": 34, "B": 10, "C": 15, "D": 11, "E": 18, "F": 15, "G": 14})
    s.etiket("A5", "Yıllık enflasyon (ÖRNEK)")
    s.girdi("C5", enf, YUZDE1)
    s.etiket("A6", "Birikimin aylık getirisi (ÖRNEK)")
    s.girdi("C6", getiri, YUZDE1)
    s.etiket("A7", "Her ay hedeflere ayırabileceğim")
    s.girdi("C7", ayirabilen, TL0)
    s.baslik_satiri(9, ["Hedef", "Öncelik", "Bugünkü tutar", "Kaç ay sonra", "O günkü tahmini tutar", "Aylık katkı", "Plana dahil mi?"])
    for i, (h, onc, tutar, ay, dahil) in enumerate(hedefler):
        r = 10 + i
        s.girdi(f"A{r}", h)
        s.girdi(f"B{r}", onc, SAYI)
        s.girdi(f"C{r}", tutar, TL0)
        s.girdi(f"D{r}", ay, SAYI)
        s.formul(f"E{r}", f"=C{r}*(1+$C$5)^(D{r}/12)", TL0)
        s.formul(f"F{r}", f"=IFERROR(-PMT($C$6,D{r},0,E{r}),0)", TL0)
        s.girdi(f"G{r}", dahil)
    son = 9 + len(hedefler)
    s.ara("A15", "Plana dahil hedeflerin aylık katkısı")
    s.formul("F15", f'=SUMIF(G10:G{son},"Evet",F10:F{son})', TL0, kalin=True)
    s.etiket("A16", "Her ay ayırabileceğim")
    s.formul("F16", "=C7", TL0)
    s.etiket("A17", "Fark")
    s.formul("F17", "=F16-F15", TL0)
    s.ara("A18", "Durum")
    s.formul("F18", '=IF(F17>=0,"PLAN YETİYOR","PLAN YETMİYOR — süre, tutar ya da öncelik değişmeli")', kalin=True)
    s.not_("A20", "Tahmini tutar = bugünkü tutar × (1 + yıllık enflasyon) ^ (ay / 12). Enflasyon ve getiri örnek varsayımdır.")
    return s


def hafta4() -> Workbook:
    wb = yeni_kitap()

    # --- 4.1 Risk ve Getiri
    s = Sayfa(wb, "Risk ve Getiri", "Risk ve Getiri · üç uydurma senaryo (4.1)", NOT_YATIRIM,
              {"A": 44, "B": 13, "C": 13, "D": 13, "E": 14, "F": 14, "G": 14})
    s.etiket("A5", "Başlangıç tutarı")
    s.girdi("C5", 10000, TL0)
    s.etiket("A6", "Yıllık enflasyon (ÖRNEK)")
    s.girdi("C6", 0.2, YUZDE1)
    s.baslik_satiri(8, ["Yıl", "Senaryo A getirisi", "Senaryo B getirisi", "Senaryo C getirisi", "A ile değer", "B ile değer", "C ile değer"])
    for i in range(10):
        r = 9 + i
        s.ws[f"A{r}"] = i + 1
        for j, k in enumerate("ABC"):
            s.girdi(f"{'BCD'[j]}{r}", SENARYO[k][i], YUZDE0)
            deger_s = "EFG"[j]
            getiri_s = "BCD"[j]
            s.formul(f"{deger_s}{r}", f"=$C$5*(1+{getiri_s}{r})" if i == 0 else f"={deger_s}{r - 1}*(1+{getiri_s}{r})", TL0)
    st = {20: ("Ortalama yıllık getiri", "AVERAGE"), 21: ("En kötü yıl", "MIN"), 22: ("En iyi yıl", "MAX"),
          23: ("Dalgalanma (standart sapma)", "STDEV.P")}
    for r, (e, fn) in st.items():
        s.etiket(f"A{r}", e)
        for c in "BCD":
            s.formul(f"{c}{r}", f"={fn}({c}9:{c}18)", YUZDE1)
    s.ara("A24", "10 yıl sonunda değer")
    s.ara("A25", "Yıllık bileşik getiri")
    s.ara("A26", "Enflasyondan arındırılmış (reel) yıllık getiri")
    for c, d in zip("BCD", "EFG"):
        s.formul(f"{c}24", f"={d}18", TL0, kalin=True)
        s.formul(f"{c}25", f"=({c}24/$C$5)^(1/10)-1", YUZDE1, kalin=True)
        s.formul(f"{c}26", f"=(1+{c}25)/(1+$C$6)-1", YUZDE1, kalin=True)
    s.not_("A28", "Ortalama yüksek olsa bile dalgalanma büyükse bileşik getiri düşebilir. Senaryolar ders için uydurulmuştur.")

    # --- 4.2 Çeşitlendirme
    s = Sayfa(wb, "Çeşitlendirme", "Çeşitlendirme · iki uydurma varlık (4.2)", NOT_YATIRIM,
              {"A": 30, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12})
    s.etiket("A6", "X'in payı →")
    for c, w in zip("DEFGH", [0, 0.25, 0.5, 0.75, 1]):
        s.girdi(f"{c}6", w, YUZDE0)
    s.baslik_satiri(7, ["Yıl", "X getirisi", "Y getirisi", "Karma 1", "Karma 2", "Karma 3", "Karma 4", "Karma 5"])
    for i in range(10):
        r = 8 + i
        s.ws[f"A{r}"] = i + 1
        s.girdi(f"B{r}", X[i], YUZDE0)
        s.girdi(f"C{r}", Y[i], YUZDE0)
        for c in "DEFGH":
            s.formul(f"{c}{r}", f"={c}$6*$B{r}+(1-{c}$6)*$C{r}", YUZDE1)
    st = {19: ("Ortalama getiri", "=AVERAGE({c}8:{c}17)", YUZDE1), 20: ("Dalgalanma (standart sapma)", "=STDEV.P({c}8:{c}17)", YUZDE1),
          21: ("En kötü yıl", "=MIN({c}8:{c}17)", YUZDE1), 22: ("Zararlı yıl sayısı", '=COUNTIF({c}8:{c}17,"<0")', SAYI)}
    for r, (e, f, b) in st.items():
        s.ara(f"A{r}", e)
        for c in "BCDEFGH":
            s.formul(f"{c}{r}", f.format(c=c), b, kalin=(r == 20))
    s.not_("A24", "Karma: X ve Y'ye üstteki paylarla bölünmüş tutar. İki varlık farklı yıllarda düştüğü için karışımın dalgalanması ikisinden de düşük.")
    s.not_("A25", "Çeşitlendirme riski azaltabilir, yok etmez: piyasanın genelini etkileyen risk kalır (SPK yatırımcı kitapçığı).")

    # --- 4.3 Hedef Planı
    hedef_sayfasi(wb, "Hedef Planı", "Finansal Hedef Planı (4.3)", NOT_KISI, HEDEFLER, 0.2, 0.01, 9000)
    return wb


# ============================================================ 4.4 Bitirme: şablon + örnek
KONTROL = [
    ("Ç1", "Dosya açılıyor, hiçbir hücrede #DIV/0!, #REF!, #VALUE! yok; girdi ve formül hücreleri ayrı"),
    ("Ç2", "Bütçe: kategori, plan, gerçekleşen, fark; toplamlar formülle"),
    ("Ç3", "12 aylık nakit akışı: açılış → giriş → çıkış → kapanış, aylar birbirine bağlı, açık aylar işaretli"),
    ("Ç4", "Borç ya da birikim: ödeme planı veya aylık katkı DEVRESEL_ÖDEME (PMT) ile; toplam faiz yazılı"),
    ("Ç5", "Enflasyon: hedef tutar enflasyon varsayımıyla güncellendi; varsayım etiketli"),
    ("Ç6", "Karar: en az 3 hedef, öncelik sırası, 'plan yetiyor mu?' sonucu ve bir değişiklik denemesi"),
    ("Ç7", "Tek sayfa rapor: durum, 3 karar, varsayımlar listesi (her oran: kaynak + tarih ya da 'ÖRNEK')"),
    ("Ç8", "KIRMIZI ÇİZGİ: gerçek kimlik/IBAN/kart no yok · ürün/hisse/fon önerisi ve getiri vaadi yok · kaynaksız 'bugünkü oran' yok"),
]


def bitirme(ornek: bool) -> Workbook:
    wb = yeni_kitap()
    not_ = NOT_KISI if ornek else "ŞABLON · kendi (ya da açıkça 'uydurma' diye işaretli) verinizle doldurun · gerçek kimlik, IBAN, kart numarası YAZMAYIN"

    s = Sayfa(wb, "Başla", "Bitirme Projesi · Kişisel Finansal Plan (4.4)", not_, {"A": 8, "B": 100, "C": 14})
    s.not_("A5", "Sırayla doldurun: Bütçe → Nakit Akışı → Hedefler → Varsayımlar. Sonra aşağıdaki listeyi işaretleyin.")
    s.baslik_satiri(7, ["Ölçüt", "Ne istenir?", "Tamam mı?"])
    for i, (k, m) in enumerate(KONTROL):
        r = 8 + i
        s.etiket(f"A{r}", k, kalin=True)
        s.etiket(f"B{r}", m)
        s.girdi(f"C{r}", "Evet" if ornek else "Hayır")
    s.ara("B17", "Tamamlanan ölçüt sayısı (8 olmalı)")
    s.formul("C17", '=COUNTIF(C8:C15,"Evet")', SAYI, kalin=True)
    s.ara("B18", "Sonuç")
    s.formul("C18", '=IF(C17=8,"TESLİME HAZIR","EKSİK VAR")', kalin=True)

    s = Sayfa(wb, "Bütçe", "Aylık Bütçe", not_, {"A": 30, "B": 12, "C": 14, "D": 14, "E": 16, "F": 12})
    s.baslik_satiri(5, ["Kategori", "Zorunlu mu?", "Plan", "Gerçekleşen", "Fark (plan − gerçek)", "Durum"])
    gercek = {"Kira": 15000, "Faturalar": 3700, "Market": 8900, "Ulaşım": 1850, "Dışarıda yemek": 1850,
              "Abonelik": 1350, "Giyim": 1800, "Sağlık": 600, "Diğer": 1550}
    for i in range(10):
        r = 6 + i
        k = KATEGORILER[i] if (ornek and i < len(KATEGORILER)) else ""
        s.girdi(f"A{r}", k or None)
        s.girdi(f"B{r}", ZORUNLU.get(k) if k else None)
        s.girdi(f"C{r}", PLAN.get(k, 0) if k else 0, TL0)
        s.girdi(f"D{r}", gercek.get(k, 0) if k else 0, TL0)
        s.formul(f"E{r}", f"=C{r}-D{r}", TL0)
        s.formul(f"F{r}", f'=IF(E{r}<0,"AŞILDI","")')
    s.ara("A16", "Toplam gider")
    for c in "CD":
        s.formul(f"{c}16", f"=SUM({c}6:{c}15)", TL0, kalin=True)
    s.formul("E16", "=C16-D16", TL0, kalin=True)
    s.etiket("A18", "Aylık gelir")
    s.girdi("C18", 45500 if ornek else 0, TL0)
    s.etiket("A19", "Birikime ayrılan")
    s.girdi("C19", 4000 if ornek else 0, TL0)
    s.ara("A20", "Kalan (gelir − gerçekleşen gider − birikim)")
    s.formul("C20", "=C18-D16-C19", TL0, kalin=True)
    s.etiket("A21", "Zorunlu giderler (gerçekleşen)")
    s.formul("C21", '=SUMIF(B6:B15,"Evet",D6:D15)', TL0)
    s.ara("A23", "Acil durum fonu")
    s.etiket("A24", "Hedef: kaç aylık zorunlu gider?")
    s.girdi("C24", 3 if ornek else 0, SAYI)
    s.etiket("A25", "Acil fon hedefi")
    s.formul("C25", "=C21*C24", TL0)
    s.etiket("A26", "Şu an ayrılmış tutar")
    s.girdi("C26", 12000 if ornek else 0, TL0)
    s.etiket("A27", "Her ay ayırabileceğim")
    s.girdi("C27", 4000 if ornek else 0, TL0)
    s.ara("A28", "Hedefe kaç ayda ulaşırım?")
    s.formul("C28", "=IF(C25-C26<=0,0,IFERROR(ROUNDUP((C25-C26)/C27,0),0))", SAYI, kalin=True)

    s = Sayfa(wb, "Nakit Akışı", "12 Aylık Nakit Akışı", not_, {"A": 32, **{c: 10.5 for c in AY_SUTUN}, "N": 12.5})
    s.baslik_satiri(5, ["Kalem", *AYLAR, "Yıl toplamı"])
    duzensiz = {"M": 15000, "L": 5000} if ornek else {}
    donemsel = {"D": 14000, "H": 25000, "J": 12000, "M": 6000} if ornek else {}
    for r, e in {6: "Açılış bakiyesi", 7: "Düzenli gelir", 8: "Düzensiz gelir", 9: "Düzenli gider",
                 10: "Dönemsel gider", 11: "Birikime ayrılan", 12: "Kapanış bakiyesi", 13: "Durum"}.items():
        (s.ara if r == 12 else s.etiket)(f"A{r}", e)
    for i, c in enumerate(AY_SUTUN):
        if i:
            s.formul(f"{c}6", f"={AY_SUTUN[i - 1]}12", TL0)
        else:
            s.girdi(f"{c}6", 2000 if ornek else 0, TL0)
        s.formul(f"{c}7", "='Bütçe'!$C$18", TL0)
        s.girdi(f"{c}8", duzensiz.get(c, 0), TL0)
        s.formul(f"{c}9", "='Bütçe'!$D$16", TL0)
        s.girdi(f"{c}10", donemsel.get(c, 0), TL0)
        s.formul(f"{c}11", "='Bütçe'!$C$19", TL0)
        s.formul(f"{c}12", f"={c}6+{c}7+{c}8-{c}9-{c}10-{c}11", TL0, kalin=True)
        s.formul(f"{c}13", f'=IF({c}12<0,"AÇIK","")')
    for r in (7, 8, 9, 10, 11):
        s.formul(f"N{r}", f"=SUM(B{r}:M{r})", TL0)
    s.etiket("A15", "En düşük kapanış")
    s.formul("B15", "=MIN(B12:M12)", TL0)
    s.etiket("A16", "Açık veren ay sayısı")
    s.formul("B16", '=COUNTIF(B12:M12,"<0")', SAYI)

    hedefler = HEDEFLER if ornek else [(None, None, 0, 12, "Hayır")] * 4
    hedef_sayfasi(wb, "Hedefler", "Finansal Hedefler", not_, hedefler, 0.2 if ornek else 0, 0.01 if ornek else 0,
                  9000 if ornek else 0)

    s = Sayfa(wb, "Varsayımlar", "Varsayımlar ve Kaynaklar", not_, {"A": 36, "B": 14, "C": 60, "D": 14})
    s.baslik_satiri(5, ["Varsayım", "Değer", "Kaynak (ya da 'ÖRNEK — kaynak yok')", "Tarih"])
    satirlar = [("Yıllık enflasyon", "='Hedefler'!C5", "ÖRNEK — resmî veri için TÜİK TÜFE bültenine bakın", "2026-09-25"),
                ("Birikimin aylık getirisi", "='Hedefler'!C6", "ÖRNEK — ürün önerisi değildir", "2026-09-25")] if ornek else \
               [("Yıllık enflasyon", "='Hedefler'!C5", None, None), ("Birikimin aylık getirisi", "='Hedefler'!C6", None, None)]
    for i, (a, v, k, t) in enumerate(satirlar):
        r = 6 + i
        s.etiket(f"A{r}", a)
        s.formul(f"B{r}", v, YUZDE1)
        s.girdi(f"C{r}", k)
        s.girdi(f"D{r}", t)
    for r in range(8, 12):
        for c in "ABCD":
            s.girdi(f"{c}{r}", None)
    return wb


KITAPLAR = {
    "hafta1-paranin-temelleri.xlsx": hafta1,
    "hafta2-borc-ve-birikim.xlsx": hafta2,
    "hafta3-isletme-finansi.xlsx": hafta3,
    "hafta4-karar-vermek.xlsx": hafta4,
    "bitirme-sablonu.xlsx": lambda: bitirme(False),
    "bitirme-ornek.xlsx": lambda: bitirme(True),
}


def main() -> None:
    CIKTI.mkdir(exist_ok=True)
    for ad, f in KITAPLAR.items():
        wb = f()
        wb.properties.creator = "Finans 101 lab (uret.py)"
        wb.properties.title = ad
        yol = CIKTI / ad
        wb.save(yol)
        print(f"{yol.relative_to(LAB)} · {len(wb.sheetnames)} sayfa: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
