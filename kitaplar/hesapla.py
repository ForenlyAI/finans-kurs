"""Çalışma kitaplarındaki formülleri GERÇEKTEN çalıştırır (formulas kütüphanesi, Excel hesap motorunun Python eşi).

Kullanım (lab klasöründen):
    uv run python kitaplar/hesapla.py [kitap.xlsx ...] [--cikti klasör]     → hesap/<kitap>.json

Çıktı: her sayfanın dolu her hücresi için {"f": formül ya da null, "v": değer}. Formül hatası (#DIV/0!, #REF! …)
varsa listelenir ve çıkış kodu 1 olur. Görüntüleyici ve grafikler değerleri BU dosyadan okur; bağımsız kontrol
(kontrol_degerler.py) aynı değerleri formül kullanmadan yeniden hesaplar.
"""

from __future__ import annotations

import json
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
import formulas  # noqa: E402
import numpy as np  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

LAB = Path(__file__).resolve().parent.parent
KITAPLAR = LAB / "calisma-kitaplari"
HESAP = LAB / "hesap"


def sade(v):
    if isinstance(v, np.ndarray):
        v = v.ravel()[0] if v.size else None
    if hasattr(v, "value") and not isinstance(v, (str, int, float)):
        v = v.value
        if isinstance(v, np.ndarray):
            v = v.ravel()[0]
    if isinstance(v, (np.floating,)):
        return float(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, np.bool_):
        return bool(v)
    if isinstance(v, (int, float, str, bool)) or v is None:
        return v
    s = str(v)
    return s if s.startswith("#") else s


def hesapla(yol: Path) -> tuple[dict, list[str]]:
    model = formulas.ExcelModel().loads(str(yol)).finish()
    cozum = model.calculate()
    anahtarlar = {k.upper(): k for k in cozum}
    wb = load_workbook(yol)  # formüller metin olarak
    sonuc: dict[str, dict] = {}
    hatalar: list[str] = []
    for ws in wb.worksheets:
        hucreler = {}
        for satir in ws.iter_rows():
            for c in satir:
                if c.value is None:
                    continue
                if isinstance(c.value, str) and c.value.startswith("="):
                    k = f"'[{yol.name}]{ws.title}'!{c.coordinate}".upper()
                    if k not in anahtarlar:
                        hatalar.append(f"{ws.title}!{c.coordinate}: hesaplanmadı")
                        v = None
                    else:
                        v = sade(cozum[anahtarlar[k]])
                    if isinstance(v, str) and v.startswith("#"):
                        hatalar.append(f"{ws.title}!{c.coordinate}: {v}")
                    hucreler[c.coordinate] = {"f": c.value, "v": v}
                else:
                    hucreler[c.coordinate] = {"f": None, "v": c.value}
        sonuc[ws.title] = hucreler
    return sonuc, hatalar


def main() -> int:
    argv = sys.argv[1:]
    hedef_klasor = HESAP
    if "--cikti" in argv:
        i = argv.index("--cikti")
        hedef_klasor = Path(argv[i + 1]).resolve()
        del argv[i:i + 2]
    hedef_klasor.mkdir(exist_ok=True)
    yollar = [Path(a).resolve() for a in argv] or sorted(KITAPLAR.glob("*.xlsx"))
    kod = 0
    for yol in yollar:
        sonuc, hatalar = hesapla(yol)
        formul = sum(1 for s in sonuc.values() for h in s.values() if h["f"])
        hedef = hedef_klasor / (yol.stem + ".json")
        hedef.write_text(json.dumps({"kitap": yol.name, "sayfalar": sonuc}, ensure_ascii=False, indent=1), encoding="utf-8")
        durum = "GEÇTİ" if not hatalar else "KALDI"
        print(f"{yol.name:34} {formul:4} formül hesaplandı · hata {len(hatalar)} · {durum}")
        for h in hatalar[:10]:
            print("   ", h)
        if hatalar:
            kod = 1
    return kod


if __name__ == "__main__":
    sys.exit(main())
