"""BAĞIMSIZ değer kontrolü: her formül hücresini formül KULLANMADAN, düz Python'la yeniden hesaplar.

Kullanım (lab klasöründen):
    uv run python kitaplar/kontrol_degerler.py [--hesap-klasoru hesap] [--kitaplar calisma-kitaplari]

Üç kapı (hepsi geçerse GEÇTİ, çıkış 0):
  1. KAPSAM   — kitaptaki HER formül hücresinin burada bir beklenen değeri var (kontrolsüz formül = KALDI)
  2. EŞİTLİK  — hesapla.py'nin (formül motoru) değeri ile bu dosyanın değeri aynı (sayı: |fark| ≤ 0,000001; metin: birebir)
  3. ÇIPA     — kitaplar/capalar.json'daki elle hesaplanmış değerler ve satır etiketleri tutuyor
Girdi hücreleri (sarı) kitaptan openpyxl ile okunur; formül metni hiç okunmaz, yalnız formül OLDUĞU bilinir.
"""

from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

LAB = Path(__file__).resolve().parent.parent
AY = [get_column_letter(i) for i in range(2, 14)]  # B..M
TOL = 1e-6          # formül motoru ↔ bağımsız model: yalnız kayan nokta farkı kabul
TOL_CAPA = 0.005    # çıpalar elle, 2 ondalığa (oranlar 6 ondalığa) yuvarlanmış yazıldı


class Kitap:
    """Girdi hücrelerini okur, beklenen değerleri toplar."""

    def __init__(self, yol: Path):
        self.wb = load_workbook(yol)
        self.beklenen: dict[tuple[str, str], object] = {}

    def g(self, sayfa: str, adres: str):
        v = self.wb[sayfa][adres].value
        assert not (isinstance(v, str) and v.startswith("=")), f"{sayfa}!{adres} girdi değil, formül"
        return 0 if v is None else v

    def gm(self, sayfa: str, adres: str):
        """girdi metni (boş → '')"""
        v = self.wb[sayfa][adres].value
        return "" if v is None else v

    def e(self, sayfa: str, adres: str, deger):
        self.beklenen[(sayfa, adres)] = deger
        return deger

    def formul_hucreleri(self) -> set[tuple[str, str]]:
        return {(ws.title, c.coordinate) for ws in self.wb.worksheets for r in ws.iter_rows() for c in r
                if isinstance(c.value, str) and c.value.startswith("=")}


def pmt(r: float, n: int, pv: float, fv: float = 0.0) -> float:
    """Excel PMT işaret kuralıyla, kapalı formül (döngüsüz)."""
    if r == 0:
        return -(pv + fv) / n
    k = (1 + r) ** n
    return -(pv * k + fv) * r / (k - 1)


def std_p(xs):
    m = sum(xs) / len(xs)
    return math.sqrt(sum((x - m) ** 2 for x in xs) / len(xs))


def roundup(x: float) -> int:
    return math.ceil(round(x, 9))


# ------------------------------------------------------------------ Hafta 1
def hafta1(k: Kitap):
    G = "Gelir-Gider"
    islem = [(k.gm(G, f"C{r}"), k.gm(G, f"D{r}"), k.g(G, f"E{r}")) for r in range(6, 26)]
    gelir = k.e(G, "H6", sum(t for _, tur, t in islem if tur == "Gelir"))
    gider = k.e(G, "H7", sum(t for _, tur, t in islem if tur == "Gider"))
    birikim = k.e(G, "H8", sum(t for _, tur, t in islem if tur == "Birikim"))
    k.e(G, "H9", gelir - gider - birikim)
    k.e(G, "H10", birikim / gelir if gelir else 0)
    kat_top = 0
    for r in range(13, 22):
        kat = k.gm(G, f"G{r}")
        v = k.e(G, f"H{r}", sum(t for c, _, t in islem if c == kat))
        kat_top += v
    k.e(G, "H22", kat_top)
    k.e(G, "H23", "EVET" if round(kat_top - gider, 2) == 0 else "HAYIR — kategori eksik")

    B = "Bütçe"
    gelir_b = k.e(B, "C19", gelir)
    k.e(B, "D19", gelir)
    top_p = top_g = zor_p = zor_g = 0
    for r in range(6, 15):
        kat = k.gm(B, f"A{r}")
        plan = k.g(B, f"C{r}")
        gercek = k.e(B, f"D{r}", sum(t for c, _, t in islem if c == kat))
        fark = k.e(B, f"E{r}", plan - gercek)
        k.e(B, f"F{r}", gercek / gelir_b if gelir_b else 0)
        k.e(B, f"G{r}", "AŞILDI" if fark < 0 else "")
        top_p += plan
        top_g += gercek
        if k.gm(B, f"B{r}") == "Evet":
            zor_p += plan
            zor_g += gercek
    k.e(B, "C15", top_p)
    k.e(B, "D15", top_g)
    k.e(B, "E15", top_p - top_g)
    k.e(B, "F15", top_g / gelir_b if gelir_b else 0)
    k.e(B, "C17", zor_p)
    k.e(B, "D17", zor_g)
    bir_plan = k.g(B, "C20")
    k.e(B, "D20", birikim)
    k.e(B, "C21", gelir_b - top_p - bir_plan)
    k.e(B, "D21", gelir_b - top_g - birikim)

    N = "Nakit Akışı"
    kapanis = []
    acilis = k.g(N, "B6")
    toplam = {r: 0 for r in (7, 8, 9, 10, 11)}
    for i, c in enumerate(AY):
        if i:
            acilis = k.e(N, f"{c}6", kapanis[-1])
        satir = {7: k.e(N, f"{c}7", gelir_b), 8: k.g(N, f"{c}8"), 9: k.e(N, f"{c}9", top_g),
                 10: k.g(N, f"{c}10"), 11: k.e(N, f"{c}11", bir_plan)}
        kap = k.e(N, f"{c}12", acilis + satir[7] + satir[8] - satir[9] - satir[10] - satir[11])
        k.e(N, f"{c}13", "AÇIK" if kap < 0 else "")
        kapanis.append(kap)
        for r in toplam:
            toplam[r] += satir[r]
    for r, v in toplam.items():
        k.e(N, f"N{r}", v)
    k.e(N, "B15", min(kapanis))
    k.e(N, "B16", sum(1 for x in kapanis if x < 0))

    A = "Acil Durum Fonu"
    zor = k.e(A, "C5", zor_g)
    hedef = k.e(A, "C7", zor * k.g(A, "C6"))
    mevcut = k.g(A, "C8")
    eksik = k.e(A, "C9", max(0, hedef - mevcut))
    aylik = k.g(A, "C10")
    k.e(A, "C11", 0 if eksik == 0 else roundup(eksik / aylik))
    bas = mevcut
    for i in range(24):
        r = 14 + i
        k.e(A, f"B{r}", bas)
        ek = k.e(A, f"C{r}", min(aylik, max(0, hedef - bas)))
        son = k.e(A, f"D{r}", bas + ek)
        k.e(A, f"E{r}", "EVET" if son >= hedef else "")
        bas = son


# ------------------------------------------------------------------ Hafta 2
def hafta2(k: Kitap):
    S = "Bileşik Faiz"
    P, r_, n = k.g(S, "C5"), k.g(S, "C6"), k.g(S, "C7")
    for y in range(11):
        row = 10 + y
        basit = k.e(S, f"B{row}", P * (1 + r_ * y))
        bilesik = P
        for _ in range(y):          # yıl yıl ekleyerek (üs almadan)
            bilesik = bilesik * (1 + r_)
        k.e(S, f"C{row}", bilesik)
        k.e(S, f"D{row}", bilesik - basit)
    k.e(S, "C22", P * (1 + r_) ** n)

    K = "Kredi Kartı"
    borc, akdi, gec, asg = k.g(K, "C5"), k.g(K, "C6"), k.g(K, "C7"), k.g(K, "C8")
    for i in range(4):
        row = 11 + i
        oran = k.e(K, f"B{row}", asg) if i == 2 else k.g(K, f"B{row}")
        odenen = k.e(K, f"C{row}", borc * oran)
        kalan = k.e(K, f"D{row}", borc - odenen)
        k.e(K, f"E{row}", "—" if kalan == 0 else ("akdi" if oran >= asg else "gecikme"))
        o = k.e(K, f"F{row}", 0 if kalan == 0 else (akdi if oran >= asg else gec))
        k.e(K, f"G{row}", kalan * o)

    L = "Kredi Planı"
    tutar, r_, n = k.g(L, "C5"), k.g(L, "C6"), k.g(L, "C7")
    taksit = k.e(L, "C8", -pmt(r_, n, tutar))
    k.e(L, "C9", taksit * n)
    k.e(L, "C10", taksit * n - tutar)
    borc = tutar
    t = {"C": 0, "D": 0, "E": 0}
    for i in range(12):
        row = 13 + i
        k.e(L, f"B{row}", borc)
        k.e(L, f"C{row}", taksit)
        faiz = k.e(L, f"D{row}", borc * r_)
        ana = k.e(L, f"E{row}", taksit - faiz)
        borc = k.e(L, f"F{row}", borc - ana)
        t["C"] += taksit
        t["D"] += faiz
        t["E"] += ana
    for c, v in t.items():
        k.e(L, f"{c}25", v)
    k.e(L, "C27", tutar * r_)   # 1. ayın faizi = anapara × oran

    H = "Birikim Hedefi"
    hedef, n, r_, mevcut = k.g(H, "C5"), k.g(H, "C6"), k.g(H, "C7"), k.g(H, "C8")
    katki = k.e(H, "C9", -pmt(r_, n, -mevcut, hedef))
    k.e(H, "C10", (hedef - mevcut) / n)
    k.e(H, "C11", hedef - mevcut - katki * n)
    bas = mevcut
    for i in range(24):
        row = 14 + i
        k.e(H, f"B{row}", bas)
        k.e(H, f"C{row}", katki)
        getiri = k.e(H, f"D{row}", bas * r_)
        bas = k.e(H, f"E{row}", bas + katki + getiri)

    E = "Enflasyon"
    sepet, enf, faiz = k.g(E, "C5"), k.g(E, "C6"), k.g(E, "C7")
    for y in range(6):
        row = 10 + y
        k.e(E, f"B{row}", sepet * (1 + enf) ** y)
        k.e(E, f"C{row}", sepet / (1 + enf) ** y)
        d = k.e(E, f"D{row}", sepet * (1 + faiz) ** y)
        k.e(E, f"E{row}", d / (1 + enf) ** y)
    k.e(E, "C17", (1 + faiz) / (1 + enf) - 1)
    yeni = k.e(E, "C18", hedef * (1 + enf) ** 2)
    k.e(E, "C19", -pmt(r_, n, -mevcut, yeni))


# ------------------------------------------------------------------ Hafta 3
def hafta3(k: Kitap):
    S = "Kâr ve Nakit"
    pesin = k.g(S, "C5")
    alti = AY[:6]
    satis = [k.g(S, f"{c}9") for c in alti]
    gider = [k.g(S, f"{c}10") for c in alti]
    top = {r: 0 for r in range(9, 16)}
    cum_k = cum_n = 0
    cum_nakitler = []
    for i, c in enumerate(alti):
        kar = k.e(S, f"{c}11", satis[i] - gider[i])
        p = k.e(S, f"{c}12", satis[i] * pesin)
        v = k.g(S, f"{c}13") if i < 2 else k.e(S, f"{c}13", satis[i - 2] * (1 - pesin))
        giris = k.e(S, f"{c}14", p + v)
        akis = k.e(S, f"{c}15", giris - gider[i])
        cum_k = k.e(S, f"{c}16", cum_k + kar)
        cum_n = k.e(S, f"{c}17", cum_n + akis)
        cum_nakitler.append(cum_n)
        k.e(S, f"{c}18", "EVET" if (cum_k > 0 and cum_n < 0) else "")
        for r, val in zip(range(9, 16), (satis[i], gider[i], kar, p, v, giris, akis)):
            top[r] += val
    for r, val in top.items():
        k.e(S, f"H{r}", val)
    k.e(S, "C20", -min(cum_nakitler))

    B = "Başabaş"
    sabit, fiyat, degisken, beklenen = k.g(B, "C5"), k.g(B, "C6"), k.g(B, "C7"), k.g(B, "C12")
    katki = k.e(B, "C8", fiyat - degisken)
    k.e(B, "C9", katki / fiyat)
    bb = k.e(B, "C10", roundup(sabit / katki))
    k.e(B, "C11", bb * fiyat)
    k.e(B, "C13", beklenen - bb)
    k.e(B, "C14", beklenen * katki - sabit)
    for i in range(11):
        row = 17 + i
        adet = k.g(B, f"A{row}")
        ciro = k.e(B, f"B{row}", adet * fiyat)
        dm = k.e(B, f"C{row}", adet * degisken)
        k.e(B, f"D{row}", sabit)
        tm = k.e(B, f"E{row}", dm + sabit)
        k.e(B, f"F{row}", ciro - tm)

    F = "Fiyatlama"
    dm = k.e(F, "C5", degisken)
    sb = k.e(F, "C6", sabit)
    ad = k.e(F, "C7", beklenen)
    pay = k.e(F, "C8", sb / ad)
    tam = k.e(F, "C9", dm + pay)
    ek = k.g(F, "C10")
    mf = k.e(F, "C11", tam * (1 + ek))
    k.e(F, "C12", (mf - tam) / mf)
    for i in range(5):
        row = 15 + i
        p = k.e(F, f"A{row}", mf) if i == 3 else k.g(F, f"A{row}")
        kp = k.e(F, f"B{row}", p - dm)
        k.e(F, f"C{row}", roundup(sb / kp) if kp != 0 else "—")
        kar = k.e(F, f"D{row}", ad * kp - sb)
        k.e(F, f"E{row}", kar / (ad * p) if ad * p else 0)

    N = "Nakit Akış Tablosu"
    bf = k.e(N, "C5", fiyat)
    bd = k.e(N, "C6", degisken)
    pes = k.g(N, "C7")
    top = {r: 0 for r in (10, 11, 14, 15, 16, 17, 18, 19)}
    kapanis = []
    onceki_satis = None
    acilis = k.g(N, "B20")
    for i, c in enumerate(AY):
        adet = k.g(N, f"{c}10")
        st = k.e(N, f"{c}11", adet * bf)
        p = k.e(N, f"{c}12", st * pes)
        g = k.g(N, f"{c}13") if i == 0 else k.e(N, f"{c}13", onceki_satis * (1 - pes))
        giris = k.e(N, f"{c}14", p + g)
        ham = k.e(N, f"{c}15", adet * bd)
        sab = k.e(N, f"{c}16", sabit)
        yat = k.g(N, f"{c}17")
        cikis = k.e(N, f"{c}18", ham + sab + yat)
        net = k.e(N, f"{c}19", giris - cikis)
        if i:
            acilis = k.e(N, f"{c}20", kapanis[-1])
        kap = k.e(N, f"{c}21", acilis + net)
        k.e(N, f"{c}22", "KASA AÇIĞI" if kap < 0 else "")
        kapanis.append(kap)
        onceki_satis = st
        for r, val in zip((10, 11, 14, 15, 16, 17, 18, 19), (adet, st, giris, ham, sab, yat, cikis, net)):
            top[r] += val
    for r, val in top.items():
        k.e(N, f"N{r}", val)
    k.e(N, "C24", min(kapanis))
    k.e(N, "C25", sum(1 for x in kapanis if x < 0))


# ------------------------------------------------------------------ Hafta 4 + bitirme
def hedefler(k: Kitap, S: str, satirlar: range):
    enf, getiri, ayir = k.g(S, "C5"), k.g(S, "C6"), k.g(S, "C7")
    dahil_top = 0
    for row in satirlar:
        tutar, ay = k.g(S, f"C{row}"), k.g(S, f"D{row}")
        gelecek = k.e(S, f"E{row}", tutar * (1 + enf) ** (ay / 12))
        katki = k.e(S, f"F{row}", -pmt(getiri, ay, 0, gelecek) if ay else 0)
        if k.gm(S, f"G{row}") == "Evet":
            dahil_top += katki
    k.e(S, "F15", dahil_top)
    k.e(S, "F16", ayir)
    fark = k.e(S, "F17", ayir - dahil_top)
    k.e(S, "F18", "PLAN YETİYOR" if fark >= 0 else "PLAN YETMİYOR — süre, tutar ya da öncelik değişmeli")


def hafta4(k: Kitap):
    R = "Risk ve Getiri"
    bas, enf = k.g(R, "C5"), k.g(R, "C6")
    for gs, ds in zip("BCD", "EFG"):
        g = [k.g(R, f"{gs}{9 + i}") for i in range(10)]
        v = bas
        for i in range(10):
            v = k.e(R, f"{ds}{9 + i}", v * (1 + g[i]))
        k.e(R, f"{gs}20", sum(g) / 10)
        k.e(R, f"{gs}21", min(g))
        k.e(R, f"{gs}22", max(g))
        k.e(R, f"{gs}23", std_p(g))
        k.e(R, f"{gs}24", v)
        cagr = k.e(R, f"{gs}25", (v / bas) ** (1 / 10) - 1)
        k.e(R, f"{gs}26", (1 + cagr) / (1 + enf) - 1)

    C = "Çeşitlendirme"
    x = [k.g(C, f"B{8 + i}") for i in range(10)]
    y = [k.g(C, f"C{8 + i}") for i in range(10)]
    seriler = {"B": x, "C": y}
    for col in "DEFGH":
        w = k.g(C, f"{col}6")
        seriler[col] = [k.e(C, f"{col}{8 + i}", w * x[i] + (1 - w) * y[i]) for i in range(10)]
    for col, s in seriler.items():
        k.e(C, f"{col}19", sum(s) / len(s))
        k.e(C, f"{col}20", std_p(s))
        k.e(C, f"{col}21", min(s))
        k.e(C, f"{col}22", sum(1 for v in s if v < 0))

    hedefler(k, "Hedef Planı", range(10, 14))


def bitirme(k: Kitap):
    S = "Başla"
    tamam = k.e(S, "C17", sum(1 for r in range(8, 16) if k.gm(S, f"C{r}") == "Evet"))
    k.e(S, "C18", "TESLİME HAZIR" if tamam == 8 else "EKSİK VAR")

    B = "Bütçe"
    tp = tg = zg = 0
    for r in range(6, 16):
        p, g = k.g(B, f"C{r}"), k.g(B, f"D{r}")
        fark = k.e(B, f"E{r}", p - g)
        k.e(B, f"F{r}", "AŞILDI" if fark < 0 else "")
        tp += p
        tg += g
        if k.gm(B, f"B{r}") == "Evet":
            zg += g
    k.e(B, "C16", tp)
    k.e(B, "D16", tg)
    k.e(B, "E16", tp - tg)
    gelir, bir = k.g(B, "C18"), k.g(B, "C19")
    k.e(B, "C20", gelir - tg - bir)
    k.e(B, "C21", zg)
    hedef = k.e(B, "C25", zg * k.g(B, "C24"))
    eksik, aylik = hedef - k.g(B, "C26"), k.g(B, "C27")
    k.e(B, "C28", 0 if eksik <= 0 else (roundup(eksik / aylik) if aylik else 0))

    N = "Nakit Akışı"
    kapanis = []
    acilis = k.g(N, "B6")
    top = {r: 0 for r in (7, 8, 9, 10, 11)}
    for i, c in enumerate(AY):
        if i:
            acilis = k.e(N, f"{c}6", kapanis[-1])
        s = {7: k.e(N, f"{c}7", gelir), 8: k.g(N, f"{c}8"), 9: k.e(N, f"{c}9", tg), 10: k.g(N, f"{c}10"),
             11: k.e(N, f"{c}11", bir)}
        kap = k.e(N, f"{c}12", acilis + s[7] + s[8] - s[9] - s[10] - s[11])
        k.e(N, f"{c}13", "AÇIK" if kap < 0 else "")
        kapanis.append(kap)
        for r in top:
            top[r] += s[r]
    for r, v in top.items():
        k.e(N, f"N{r}", v)
    k.e(N, "B15", min(kapanis))
    k.e(N, "B16", sum(1 for v in kapanis if v < 0))

    hedefler(k, "Hedefler", range(10, 14))
    k.e("Varsayımlar", "B6", k.g("Hedefler", "C5"))
    k.e("Varsayımlar", "B7", k.g("Hedefler", "C6"))


MODELLER = {
    "hafta1-paranin-temelleri.xlsx": hafta1,
    "hafta2-borc-ve-birikim.xlsx": hafta2,
    "hafta3-isletme-finansi.xlsx": hafta3,
    "hafta4-karar-vermek.xlsx": hafta4,
    "bitirme-sablonu.xlsx": bitirme,
    "bitirme-ornek.xlsx": bitirme,
}


def esit(a, b, tol: float = TOL) -> bool:
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= tol + 1e-9 * abs(b)
    return a == b


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--kitaplar", default=str(LAB / "calisma-kitaplari"))
    ap.add_argument("--hesap-klasoru", default=str(LAB / "hesap"))
    ap.add_argument("--capalar", default=str(LAB / "kitaplar" / "capalar.json"))
    ap.add_argument("--sessiz", action="store_true")
    a = ap.parse_args()
    kdir, hdir = Path(a.kitaplar), Path(a.hesap_klasoru)
    capalar = json.loads(Path(a.capalar).read_text(encoding="utf-8"))
    genel = 0
    toplam_hucre = 0
    for ad, model in MODELLER.items():
        yol = kdir / ad
        if not yol.exists():
            print(f"{ad:34} YOK · KALDI")
            genel = 1
            continue
        k = Kitap(yol)
        model(k)
        hesap = json.loads((hdir / (yol.stem + ".json")).read_text(encoding="utf-8"))["sayfalar"]
        formuller = k.formul_hucreleri()
        eksik = sorted(formuller - set(k.beklenen))
        fazla = sorted(set(k.beklenen) - formuller)
        farklar = []
        for (s, adr), bek in k.beklenen.items():
            if (s, adr) not in formuller:
                continue
            gercek = hesap.get(s, {}).get(adr, {}).get("v")
            if not esit(gercek, bek):
                farklar.append(f"{s}!{adr}: formül motoru {gercek!r} · bağımsız {bek!r}")
        cap_hata = []
        for c in capalar.get(ad, []):
            h = hesap.get(c["sayfa"], {})
            v = h.get(c["hucre"], {}).get("v")
            et = h.get(c["etiket_hucre"], {}).get("v") if c.get("etiket_hucre") else None
            tol = c.get("tol", TOL_CAPA if isinstance(c["deger"], (int, float)) and abs(c["deger"]) >= 1 else 5e-6)
            if not esit(v, c["deger"], tol):
                cap_hata.append(f"ÇIPA {c['sayfa']}!{c['hucre']}: {v!r} ≠ {c['deger']!r} ({c['neden']})")
            if c.get("etiket") and et != c["etiket"]:
                cap_hata.append(f"ÇIPA etiket {c['sayfa']}!{c['etiket_hucre']}: {et!r} ≠ {c['etiket']!r}")
        ok = not (eksik or fazla or farklar or cap_hata)
        toplam_hucre += len(formuller)
        print(f"{ad:34} {len(formuller):4} formül · kontrolsüz {len(eksik)} · fark {len(farklar)} · "
              f"çıpa {len(capalar.get(ad, [])) - len(cap_hata)}/{len(capalar.get(ad, []))} · {'GEÇTİ' if ok else 'KALDI'}")
        if not a.sessiz:
            for m in (eksik[:5] and [f"kontrolsüz formül: {s}!{x}" for s, x in eksik[:5]]) or []:
                print("   ", m)
            for m in [f"beklenen ama formül değil: {s}!{x}" for s, x in fazla[:5]] + farklar[:8] + cap_hata:
                print("   ", m)
        if not ok:
            genel = 1
    print(f"toplam {toplam_hucre} formül hücresi bağımsız olarak yeniden hesaplandı · {'GEÇTİ' if genel == 0 else 'KALDI'}")
    return genel


if __name__ == "__main__":
    sys.exit(main())
