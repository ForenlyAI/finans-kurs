"""NEGATİF TEST: kontrol gerçekten hata yakalıyor mu?

Her bozma ayrı bir geçici kopyada yapılır → hesapla → kontrol_degerler. Kontrol KALDI demeli.
Bir bozma GEÇTİ ile sonuçlanırsa kontrol kördür ve bu betik KALDI döner.

    uv run python kitaplar/negatif_test.py
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

LAB = Path(__file__).resolve().parent.parent
BOZMALAR = [
    ("hafta2-borc-ve-birikim.xlsx", "Kredi Planı", "C8", "=-PMT(C6/12,C7,C5)", "aylık oranı bir daha 12'ye bölmek"),
    ("hafta1-paranin-temelleri.xlsx", "Nakit Akışı", "D6", "=B12", "Mart açılışını Ocak kapanışına bağlamak"),
    ("hafta3-isletme-finansi.xlsx", "Başabaş", "C10", "=ROUNDUP(C5/C6,0)", "sabit gideri katkı payı yerine fiyata bölmek"),
    ("hafta2-borc-ve-birikim.xlsx", "Kredi Planı", "C9", 60277, "formülü silip elle sayı yazmak"),
    ("hafta2-borc-ve-birikim.xlsx", "Kredi Kartı", ("C6", "C7"), None, "iki girdiyi yer değiştirmek (permütasyon — yalnız çıpa yakalar)"),
    ("hafta4-karar-vermek.xlsx", "Çeşitlendirme", "F20", "=STDEV.P(F8:F16)", "aralıktan bir yılı düşürmek"),
]


def main() -> int:
    kor = 0
    for kitap, sayfa, hucre, yeni, neden in BOZMALAR:
        with tempfile.TemporaryDirectory() as gd:
            gd = Path(gd)
            kd, hd = gd / "kitaplar", gd / "hesap"
            kd.mkdir()
            hd.mkdir()
            for x in (LAB / "calisma-kitaplari").glob("*.xlsx"):
                shutil.copy(x, kd / x.name)
            wb = load_workbook(kd / kitap)
            ws = wb[sayfa]
            if isinstance(hucre, tuple):
                a, b = hucre
                ws[a].value, ws[b].value = ws[b].value, ws[a].value
            else:
                ws[hucre].value = yeni
            wb.save(kd / kitap)
            subprocess.run([sys.executable, str(LAB / "kitaplar" / "hesapla.py"), *map(str, sorted(kd.glob("*.xlsx"))),
                            "--cikti", str(hd)], capture_output=True, text=True)
            k = subprocess.run([sys.executable, str(LAB / "kitaplar" / "kontrol_degerler.py"), "--kitaplar", str(kd),
                                "--hesap-klasoru", str(hd), "--sessiz"], capture_output=True, text=True)
            yakalandi = k.returncode != 0
            etiket = f"{sayfa}!{hucre if isinstance(hucre, str) else '↔'.join(hucre)}"
            print(f"{'YAKALANDI' if yakalandi else 'KAÇTI    '} · {etiket:26} · {neden}")
            if not yakalandi:
                kor = 1
    print("NEGATİF TEST:", "GEÇTİ (her bozma yakalandı)" if kor == 0 else "KALDI (kontrol kör)")
    return kor


if __name__ == "__main__":
    sys.exit(main())
